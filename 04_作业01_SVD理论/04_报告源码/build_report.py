#!/usr/bin/env python3
"""Build the template-based English DOCX with editable Office equations."""
from pathlib import Path
from collections import Counter
import json, re, sys
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from content import pages, S, ROOT

cache_path = Path(__file__).with_name('equations.json')
stored = json.loads(cache_path.read_text()) if cache_path.exists() else {}
cache = stored.get('equations', {}) if stored.get('version') == 4 else {}
legacy_dependencies = ROOT.parent / '01_实验01_SVD及其应用/05_构建工具'
if legacy_dependencies.exists():
    sys.path.insert(0, str(legacy_dependencies))

def office_math(latex):
    if latex not in cache:
        from latex2mathml.converter import convert
        stylesheet = Path('/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl')
        transform = etree.XSLT(etree.parse(str(stylesheet)))
        # Group parentheses before conversion, including bases of superscripts.
        # The Office stylesheet can otherwise mistake a final period for a fence.
        grouped_latex = latex.replace(r'\left(', '(').replace(r'\right)', ')')
        grouped_latex = grouped_latex.replace('(', r'\left(').replace(')', r'\right)')
        ml = etree.fromstring(convert(grouped_latex).encode())
        ns = '{http://www.w3.org/1998/Math/MathML}'
        pairs = {'(': ')', '[': ']', '{': '}', '‖': '‖', '|': '|'}
        for row in list(ml.iter(ns+'mrow'))[::-1]:
            if len(row) >= 2 and row[0].tag == ns+'mo' and row[-1].tag == ns+'mo':
                if pairs.get(row[0].text) == row[-1].text and (
                    row[0].get('stretchy') == 'true' or row[0].get('fence') == 'true'):
                    fenced = etree.Element(ns+'mfenced', open=row[0].text,
                                           close=row[-1].text, separators='')
                    inner = etree.SubElement(fenced, ns+'mrow')
                    for child in list(row)[1:-1]: inner.append(child)
                    row.getparent().replace(row, fenced)
        # pmatrix emits ungrouped parentheses around an mtable. Convert each
        # matrix's pair into one native delimiter that grows to the matrix height.
        for matrix in list(ml.iter(ns+'mtable')):
            parent = matrix.getparent()
            pos = parent.index(matrix)
            if pos > 0 and pos+1 < len(parent):
                opening, closing = parent[pos-1], parent[pos+1]
                if opening.tag == ns+'mo' and closing.tag == ns+'mo' and pairs.get(opening.text) == closing.text:
                    fenced = etree.Element(ns+'mfenced', open=opening.text, close=closing.text, separators='')
                    parent.insert(pos-1, fenced)
                    fenced.append(matrix)
                    parent.remove(opening); parent.remove(closing)
        result = transform(ml)
        tokens = lambda txt: Counter(re.findall(r'[A-Za-z0-9\u0370-\u03ff−]', txt))
        before = tokens(''.join(ml.itertext()))
        after = tokens(''.join(x.text or '' for x in result.getroot().iter(qn('m:t'))))
        if before != after:
            raise ValueError('Equation conversion altered semantic tokens: '+latex)
        cache[latex] = etree.tostring(result.getroot(), encoding='unicode')
    node = etree.fromstring(cache[latex].encode())
    for mr in node.iter(qn('m:r')):
        rp = mr.find(qn('w:rPr'))
        if rp is None:
            rp = OxmlElement('w:rPr'); mr.insert(0, rp)
        size = OxmlElement('w:sz'); size.set(qn('w:val'), '21'); rp.append(size)
    return node

doc = Document(ROOT/'01_指导资料/Original_Report_Template.docx')
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'; normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.06
for name in ['Heading 1', 'Heading 2']:
    if name not in doc.styles: doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style = doc.styles[name]
    style.base_style = normal
    style.font.name = 'Times New Roman'
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string('17365D')
    style.font.size = Pt(15 if name == 'Heading 1' else 11.5)
    style.paragraph_format.space_before = Pt(5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1
    style.paragraph_format.keep_with_next = True
doc.core_properties.title = 'Singular Value Decomposition - Homework Solutions'
doc.core_properties.subject = 'Theory, Algorithms, and Applications'
doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
doc.paragraphs[0].text = 'Beihang University\nSchool of Mathematical Sciences\nHomework Report'
doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in doc.paragraphs[0].runs:
    run.font.name = 'Times New Roman'; run.font.size = Pt(15); run.bold = True
doc.paragraphs[1].text = 'Complete written solutions and computational results continue on the following pages.'
for run in doc.paragraphs[1].runs:
    run.font.name = 'Times New Roman'; run.font.size = Pt(9); run.bold = False
T = doc.tables[0]
values = {
 (0,0): 'Course: __________',
 (0,2): 'Assignment: Singular Value Decomposition',
 (1,0): 'Type: Demonstration [ ]  Verification [ ]\nComprehensive [ ]  Design [ ]',
 (2,0): 'Class: 230116', (2,1): 'Name: yule yang', (2,3): 'Student ID: 23376268',
 (3,0): 'Date: __________', (3,1): 'Instructor: __________', (3,3): 'Grade: __________',
 (4,0): 'Computational environment:\nPython '+S['environment']['python']+'; NumPy '+S['environment']['numpy']+'; Matplotlib '+S['environment']['matplotlib']+'; Pillow '+S['environment']['pillow']+'. macOS on Apple Silicon. Numerical work was executed locally.',
 (5,0): 'Objectives and content:\nSolve Problems 1-15 on SVD existence and geometry, fundamental subspaces, norms, the pseudoinverse, least squares, optimal low-rank approximation, compression, PCA and regularization. Complete both prescribed numerical experiments and the optional complex-SVD challenge.',
 (6,0): 'Procedure and results:\nProvide proofs and direct calculations in the assignment task order. Compress a 512-by-512 grayscale image at ranks 5, 10, 20, 50 and 100. Generate a 100-by-80 rank-five signal plus Gaussian noise; evaluate ranks 1-20. Include error tables, four figures and the complete Python source. All '+str(S['checks_total'])+' numerical assertions passed. Detailed solutions follow.',
 (7,0): 'Analysis and summary:\nSVD connects geometry, stable estimation and optimal approximation. Distinguish exact rank from approximate rank, scalar-storage fraction from file-size reduction, and a better fit to noisy data from better recovery of a clean signal. Instructor and grade remain blank.'
}
for (i,j), txt in values.items():
    cell = T.cell(i,j); cell.text = txt
    for para in cell.paragraphs:
        para.paragraph_format.space_before = Pt(1)
        para.paragraph_format.space_after = Pt(4)
        para.paragraph_format.line_spacing = 1.02
        for run in para.runs: run.font.name = 'Times New Roman'; run.font.size = Pt(10)
for row in T.rows:
    pr = row._tr.get_or_add_trPr()
    for element in list(pr):
        if element.tag in [qn('w:trHeight'), qn('w:cantSplit')]: pr.remove(element)
    pr.append(OxmlElement('w:cantSplit'))

section = doc.add_section(WD_SECTION_START.NEW_PAGE)
section.top_margin = Mm(19); section.bottom_margin = Mm(19)
section.left_margin = Mm(20); section.right_margin = Mm(20)
section.header_distance = Mm(8); section.footer_distance = Mm(8)
section.header.is_linked_to_previous = False
hp = section.header.paragraphs[0]
hp.text = 'Singular Value Decomposition | Homework Solutions'
hp.runs[0].font.name = 'Times New Roman'; hp.runs[0].font.size = Pt(9)
section.footer.is_linked_to_previous = False
fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); fp._p.append(field)
for sec in doc.sections:
    for element in list(sec._sectPr):
        if element.tag == qn('w:docGrid'): sec._sectPr.remove(element)
for element in doc.settings._element.iter(qn('w:compatSetting')):
    if element.get(qn('w:name')) == 'compatibilityMode': element.set(qn('w:val'), '15')

for idx, pg in enumerate(pages):
    title = doc.add_paragraph(pg['title'], 'Heading 1')
    title.paragraph_format.page_break_before = idx > 0
    title.paragraph_format.space_after = Pt(9)
    for b in pg['blocks']:
        kind = b['kind']
        if kind == 'paragraph':
            para = doc.add_paragraph(b['text'])
            para.paragraph_format.widow_control = True
        elif kind == 'heading': doc.add_paragraph(b['text'], 'Heading 2')
        elif kind == 'equation':
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(2)
            para.paragraph_format.space_after = Pt(7)
            para._p.append(office_math(b['latex']))
        elif kind == 'table':
            table = doc.add_table(rows=1, cols=len(b['headers'])); table.style = 'Table Grid'
            for c, txt in zip(table.rows[0].cells, b['headers']): c.text = txt
            table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for row in b['rows']:
                for c, txt in zip(table.add_row().cells, row): c.text = str(txt)
            for i, row in enumerate(table.rows):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_before = Pt(2)
                        para.paragraph_format.space_after = Pt(2)
                        para.paragraph_format.line_spacing = 1
                        for run in para.runs:
                            run.font.name = 'Times New Roman'; run.font.size = Pt(9)
                            run.bold = i == 0
                    if i == 0:
                        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E8EDF4')
                        cell._tc.get_or_add_tcPr().append(shade)
            para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.space_before = Pt(0); para.paragraph_format.line_spacing = Pt(2)
            para.add_run().font.size = Pt(2)
        elif kind == 'figure':
            path = ROOT/'02_实验本体/results'/b['name']
            w, h = Image.open(path).size
            height = min(b['height'], 6.65*h/w)
            para = doc.add_paragraph(); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(path), height=Inches(height))
            caption = doc.add_paragraph(b['caption'])
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in caption.runs: run.font.size = Pt(9); run.italic = True
        elif kind == 'code':
            para = doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = Pt(10.8)
            para.paragraph_format.widow_control = False
            for i, line in enumerate(b['text'].splitlines()):
                if i: para.add_run().add_break()
                run = para.add_run(line); run.font.name = 'Courier New'; run.font.size = Pt(8)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F4F5F7')
            para._p.get_or_add_pPr().append(shade)
out = ROOT/'03_最终实验报告/SVD_Homework_English.docx'
doc.save(out)
cache_path.write_text(json.dumps(dict(version=4, equations=cache), ensure_ascii=False, indent=2))
(ROOT/'04_报告源码/report_content.json').write_text(json.dumps(pages, ensure_ascii=False, indent=2))
print('Saved', out, 'with', len(cache), 'editable equation definitions and', len(pages), 'content sections.')
