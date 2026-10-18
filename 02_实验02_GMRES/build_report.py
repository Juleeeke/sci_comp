"""Shared content -> LaTeX first, then editable Word using the retained template.
Run python build_report.py --tex, compile 04_LaTeX源码/main.tex, then --word.
"""
from pathlib import Path
import argparse,json,re,sys,shutil,copy
from collections import Counter
from lxml import etree
from docx import Document
from docx.shared import Pt,Mm,Inches,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
C=json.loads((ROOT/'content.json').read_text())
CACHE=ROOT/'equations.json'
cache=json.loads(CACHE.read_text()) if CACHE.exists() else {}
sys.path.insert(0,str(ROOT/'05_构建工具'))
def math(latex):
    if latex not in cache:
        from latex2mathml.converter import convert
        xsl=Path('/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl')
        if not xsl.exists():raise RuntimeError('Equation cache missing; regenerate with latex2mathml and the Microsoft Office MathML stylesheet available.')
        normalized=re.sub(r'\\\|(.*?)\\\|',lambda m:r'\left\|'+m.group(1)+r'\right\|',latex.replace(r'\left\|',r'\|').replace(r'\right\|',r'\|'))
        normalized=normalized.replace(r'\arg\min',r'\operatorname{argmin}')
        ml=etree.fromstring(convert(normalized).encode());ns='{http://www.w3.org/1998/Math/MathML}'
        # The converter emits bare delimiters around pmatrix tables. Group them
        # explicitly so Office receives a growing delimiter, not text parentheses.
        for mt in list(ml.iter(ns+'mtable')):
            parent=mt.getparent();i=parent.index(mt)
            if i and i+1<len(parent):
                left,right=parent[i-1],parent[i+1]
                if left.tag==right.tag==ns+'mo' and (left.text,right.text) in [('(',')'),('[',']')]:
                    fenced=etree.Element(ns+'mfenced',open=left.text,close=right.text,separators='')
                    inner=etree.SubElement(fenced,ns+'mrow');inner.append(mt)
                    parent.remove(left);parent.remove(right);parent.insert(i-1,fenced)
        for mover in ml.iter(ns+'mover'):
            mover.set('accent','true')
            if mover[-1].text=='^':mover[-1].text='\u0302'
            if mover[-1].text=='\u2015':mover[-1].text='\u0305'
        pairs={'(':')','[':']','{':'}','‖':'‖','|':'|'}
        for row in list(ml.iter(ns+'mrow'))[::-1]:
            if len(row)>=2 and row[0].tag==ns+'mo' and row[-1].tag==ns+'mo' and pairs.get(row[0].text)==row[-1].text and (row[0].get('stretchy')=='true' or row[0].get('fence')=='true'):
                fenced=etree.Element(ns+'mfenced',open=row[0].text,close=row[-1].text,separators='');inner=etree.SubElement(fenced,ns+'mrow')
                for child in list(row)[1:-1]:inner.append(child)
                row.getparent().replace(row,fenced)
        out=etree.XSLT(etree.parse(str(xsl)))(ml).getroot()
        tokens=lambda s:Counter(re.findall(r'[A-Za-z0-9\u0370-\u03ff−]',s))
        before=tokens(''.join(ml.itertext()));after=tokens(''.join(x.text or '' for x in out.iter(qn('m:t'))))
        if before!=after:raise ValueError('Semantic token mismatch in '+latex+' '+str(before-after)+' '+str(after-before))
        if any(re.search(r'\\[A-Za-z]+',x.text or '') for x in out.iter(qn('m:t'))):raise ValueError('Unconverted math command in '+latex)
        cache[latex]=etree.tostring(out,encoding='unicode')
    return etree.fromstring(cache[latex].encode())
def esc(s):
    m={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(m.get(x,x) for x in str(s))
def write_tex():
    (ROOT/'04_LaTeX源码/figures').mkdir(exist_ok=True)
    tex=[r'''\documentclass[a4paper,11pt]{article}
\usepackage[margin=20mm]{geometry}
\usepackage{amsmath,amssymb,graphicx,booktabs,array,float,hyperref}
\hypersetup{hidelinks}
\setlength{\parindent}{0pt}
\setlength{\parskip}{5pt}
\renewcommand{\arraystretch}{1.14}
\setlength{\tabcolsep}{4pt}
\begin{document}
\thispagestyle{empty}
\begin{center}{\Large\bfseries Beihang University\\School of Mathematical Sciences\\Laboratory Report}\end{center}
{\small Detailed derivations, results and analysis continue on the following pages.}\par
\vspace{3mm}
\begin{tabular}{|p{.22\linewidth}|p{.22\linewidth}|p{.22\linewidth}|p{.22\linewidth}|}\hline
\multicolumn{2}{|p{.47\linewidth}|}{Course: Numerical Analysis Computer Laboratory} & \multicolumn{2}{p{.43\linewidth}|}{'''+esc(C['title'])+r'''}\\\hline
\multicolumn{4}{|l|}{Experiment type: Demonstration [ ] Verification [ ] Comprehensive [ ] Design [ ]}\\\hline
Class: 230116 & \multicolumn{2}{l|}{Name: yule yang} & Student ID: 23376268\\\hline
Lab date: 2026.10.17 & \multicolumn{2}{l|}{Instructor:} & Grade:\\\hline
''']
    for title,key in [('Experimental environment','environment'),('Objectives and content','objectives'),('Procedure and results','summary'),('Analysis and conclusions','conclusions')]:
        tex.append(r'\multicolumn{4}{|p{.95\linewidth}|}{\textbf{'+esc(title)+r'}\par '+esc(C[key])+r'\par\vspace{4mm}}\\\hline')
    tex.append(r'\end{tabular}')
    for pg in C['pages']:
        tex.extend([r'\clearpage',r'\section*{'+esc(pg['title'])+'}'])
        for b in pg['blocks']:
            k=b['kind']
            if k=='p':tex.append(esc(b['text'])+'\n')
            elif k=='h':tex.append(r'\subsection*{'+esc(b['text'])+'}')
            elif k=='eq':tex.append(r'\['+b['latex']+r'\]')
            elif k=='table':
                n=len(b['headers']);col=''.join('p{'+str((1-.018*n)/n)+r'\linewidth}' for _ in range(n));tex.append(r'{\small\begin{tabular}{'+col+r'}\toprule')
                tex.append(' & '.join(esc(x) for x in b['headers'])+r'\\\midrule')
                tex.extend(' & '.join(esc(x) for x in row)+r'\\' for row in b['rows']);tex.append(r'\bottomrule\end{tabular}}\par')
            elif k=='fig':
                src=ROOT/b['path'];dst=ROOT/'04_LaTeX源码/figures'/src.name;shutil.copy2(src,dst)
                tex.append(r'\begin{figure}[H]\centering\includegraphics[width=\linewidth,height='+str(b.get('height',2.35)*25.4)+r'mm,keepaspectratio]{figures/'+dst.name+r'}\caption{'+esc(b['caption'])+r'}\end{figure}')
    tex.append(r'\end{document}');(ROOT/'04_LaTeX源码/main.tex').write_text('\n'.join(tex))
def write_word():
    if not (ROOT/'04_LaTeX源码/main.pdf').is_file():raise RuntimeError('Compile LaTeX successfully before producing Word.')
    d=Document(ROOT/'01_指导资料/Original_Report_Template.docx');d.core_properties.title=C['title'];d.core_properties.author='';d.core_properties.last_modified_by=''
    normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.line_spacing=1.03
    for name,size in [('Title',15),('Heading 1',15),('Heading 2',12)]:
        if name not in d.styles:d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_before=Pt(4);st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.keep_with_next=True
    for s in d.sections:
        s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=Mm(18);s.bottom_margin=Mm(18);s.left_margin=Mm(20);s.right_margin=Mm(20)
    for el in list(d.sections[0]._sectPr):
        if el.tag==qn('w:docGrid'):d.sections[0]._sectPr.remove(el)
    d.paragraphs[0].text='Beihang University\nSchool of Mathematical Sciences\nLaboratory Report';d.paragraphs[0].style=d.styles['Title'];d.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    d.paragraphs[1].text='Detailed derivations, results and analysis continue on the following pages.';d.paragraphs[1].style=normal
    d.paragraphs[1].runs[0].font.size=Pt(9)
    table=d.tables[0]
    values={(0,0):'Course: Numerical Analysis\nComputer Laboratory',(0,2):C['title'],(1,0):'Experiment type: Demonstration [ ] Verification [ ]\nComprehensive [ ] Design [ ]',(2,0):'Class: 230116',(2,1):'Name: yule yang',(2,3):'Student ID: 23376268',(3,0):'Lab date: 2026.10.17',(3,1):'Instructor:',(3,3):'Grade:',(4,0):'Experimental environment:\n'+C['environment'],(5,0):'Objectives and content:\n'+C['objectives'],(6,0):'Procedure and results:\n'+C['summary'],(7,0):'Analysis and conclusions:\n'+C['conclusions']}
    for (i,j),text in values.items():
        cell=table.cell(i,j);cell.text=text
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1
            for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
    for row in table.rows:
        pr=row._tr.get_or_add_trPr()
        for el in list(pr):
            if el.tag in [qn('w:trHeight'),qn('w:cantSplit')]:pr.remove(el)
        pr.append(OxmlElement('w:cantSplit'))
    d.add_section(WD_SECTION_START.NEW_PAGE)
    for s in d.sections:
        for el in list(s._sectPr):
            if el.tag==qn('w:docGrid'):s._sectPr.remove(el)
    for el in d.settings._element.iter(qn('w:compatSetting')):
        if el.get(qn('w:name'))=='compatibilityMode':el.set(qn('w:val'),'15')
    for i,pg in enumerate(C['pages']):
        ph=d.add_paragraph(pg['title'],'Heading 1');ph.paragraph_format.page_break_before=i>0
        for b in pg['blocks']:
            k=b['kind']
            if k=='p':d.add_paragraph(b['text'])
            elif k=='h':d.add_paragraph(b['text'],'Heading 2')
            elif k=='eq':
                p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p._p.append(math(b['latex']))
            elif k=='table':
                t=d.add_table(rows=1,cols=len(b['headers']));t.style='Table Grid'
                for cell,text in zip(t.rows[0].cells,b['headers']):cell.text=str(text)
                t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for row in b['rows']:
                    for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
                for j,row in enumerate(t.rows):
                    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
                            for r in p.runs:r.font.size=Pt(b.get('size',9));r.bold=j==0
                d.add_paragraph().paragraph_format.space_after=Pt(0)
            elif k=='fig':
                from PIL import Image
                w,h=Image.open(ROOT/b['path']).size
                width=min(6.65,b.get('height',2.35)*w/h)
                p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(ROOT/b['path']),width=Inches(width))
                cp=d.add_paragraph(b['caption']);cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in cp.runs:r.font.size=Pt(9);r.italic=True
    for p in d.paragraphs:
        for r in p.runs:
            if not r.font.name:r.font.name='Times New Roman'
    d.save(ROOT/'03_最终实验报告/Report.docx');CACHE.write_text(json.dumps(cache,indent=2,ensure_ascii=False))
    shutil.copy2(ROOT/'04_LaTeX源码/main.pdf',ROOT/'03_最终实验报告/Report.pdf')
    print('Created editable Word with',len(cache),'cached native equations')
a=argparse.ArgumentParser();a.add_argument('--tex',action='store_true');a.add_argument('--word',action='store_true');args=a.parse_args()
if args.tex:write_tex()
if args.word:write_word()
