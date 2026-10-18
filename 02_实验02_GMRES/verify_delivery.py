"""Audit final saved artifacts and recorded numerical arrays independently."""
from pathlib import Path
import json,hashlib,re,zipfile
import numpy as np
import pymupdf as fitz
from lxml import etree
from docx import Document
ROOT=Path(__file__).resolve().parent
C=json.loads((ROOT/'content.json').read_text());checks=[];details={}
def check(name,condition,detail=''):checks.append(dict(name=name,passed=bool(condition),detail=detail))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
S=json.loads((ROOT/'01_指导资料/source_manifest.json').read_text())
check('Every preserved source hash matches',all(sha(ROOT/'01_指导资料'/n)==v for n,v in S.items()))
for name in ['03_最终实验报告/Report.docx','03_最终实验报告/Report.pdf','04_LaTeX源码/main.tex','04_LaTeX源码/compile.sh','equations.json','README.md','COVERAGE.md','requirements.txt']:
 check('Required artifact '+name,(ROOT/name).is_file() and (ROOT/name).stat().st_size>0)
check('Compiled PDF is exact delivered PDF',sha(ROOT/'04_LaTeX源码/main.pdf')==sha(ROOT/'03_最终实验报告/Report.pdf'))
check('LaTeX log has no error or overfull box',not re.search(r'(^!|Overfull|Missing character)',(ROOT/'04_LaTeX源码/main.log').read_text(),re.M))
d=Document(ROOT/'03_最终实验报告/Report.docx');t=d.tables[0]
check('Template table dimensions retained',len(t.rows)==8 and len(t.columns)==4)
T=Document(ROOT/'01_指导资料/Original_Report_Template.docx').tables[0]
merge=lambda table:[[c._tc.get('x') or (c._tc.tcPr.xml if c._tc.tcPr is not None else '') for c in r.cells] for r in table.rows]
# Compare the actual merge metadata independent of presentation properties.
def merges(table):return [[(c._tc.tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan').get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val') if c._tc.tcPr.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}gridSpan') is not None else '1') for c in r.cells] for r in table.rows]
check('Template merged cell structure retained',merges(t)==merges(T))
check('Identity and date match user screenshot and unprovided assessment fields remain blank',all(t.cell(i,j).text==label for i,j,label in [(2,0,'Class: 230116'),(2,1,'Name: yule yang'),(2,3,'Student ID: 23376268'),(3,0,'Lab date: 2026.10.17'),(3,1,'Instructor:'),(3,3,'Grade:')]))
alltext='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t0 in d.tables for r in t0.rows for c in r.cells)
check('Authored Word text entirely English',not re.search(r'[\u4e00-\u9fff]',alltext))
check('Every assigned section present',all(pg['title'] in alltext for pg in C['pages']))
ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(ROOT/'03_最终实验报告/Report.docx') as z:
 check('DOCX ZIP integrity',z.testzip() is None);xml=etree.fromstring(z.read('word/document.xml'));eqs=xml.findall('.//m:oMath',ns)
 expected=sum(b['kind']=='eq' for pg in C['pages'] for b in pg['blocks']);check('Every formula is a native Office equation',len(eqs)==expected,f'{len(eqs)} equations')
 check('No unconverted LaTeX commands in equations',all(not re.search(r'\\[A-Za-z]+',x.text or '') for x in xml.findall('.//m:t',ns)))
 expected_figs=sum(b['kind']=='fig' for pg in C['pages'] for b in pg['blocks']);check('All figures embedded without whole page images',len(xml.findall('.//w:drawing',ns))==expected_figs,f'{expected_figs} figures')
 details['native_equations']=len(eqs);details['embedded_figures']=expected_figs
for mode,pdf in [('latex',ROOT/'03_最终实验报告/Report.pdf'),('word',ROOT/'06_验证记录/word_render/Report.pdf')]:
 if not pdf.exists():
  details[mode+'_rendering']='Not rerun locally; see recorded delivery audit and review record.'
  continue
 doc=fitz.open(pdf);texts=[p.get_text() for p in doc];details[mode+'_pages']=len(doc)
 check(mode+' correct page count',len(doc)==len(C['pages'])+1)
 check(mode+' no empty pages',all(len(t0)>150 for t0 in texts))
 check(mode+' all section titles survive rendering',all(re.sub(r'\s+','',pg['title']) in re.sub(r'\s+','',''.join(texts)) for pg in C['pages']))
 check(mode+' no replacement glyph',all('\ufffd' not in t0 and '¿' not in t0 for t0 in texts))
 check(mode+' authored text within page bounds',all(24<=s['bbox'][0] and s['bbox'][2]<=page.rect.width-20 and s['bbox'][3]<=page.rect.height-20 for page in doc for b in page.get_text('dict')['blocks'] if b['type']==0 for line in b['lines'] for s in line['spans']))
 if (ROOT/'06_验证记录'/mode).exists():check(mode+' every page rendered',all((ROOT/'06_验证记录'/mode/f'page_{i+1}.png').exists() for i in range(len(doc))))
V=json.loads((ROOT/'02_实验本体/results/verification.json').read_text());check('All numerical verification checks pass',V['passed']==V['total'],f"{V['passed']}/{V['total']}");details['numerical_verification']=f"{V['passed']}/{V['total']}"
if ROOT.name.startswith('02_实验02'):
 from scipy.sparse import load_npz
 from scipy.sparse.linalg import spsolve
 R=json.loads((ROOT/'02_实验本体/results/results.json').read_text())
 for case in R['cases']:
  tag=case['name'].replace(' ','_').replace('=','_');A=load_npz(ROOT/'02_实验本体/data'/(tag+'_A.npz'));v=np.load(ROOT/'02_实验本体/data'/(tag+'_vectors.npz'));b=v['b'];truth=v['truth'];direct=spsolve(A,b)
  if 'grid_n' not in case:check(tag+' RHS and exact solution match',np.array_equal(b,A@np.ones(A.shape[0])))
  else:
   N=case['grid_n'];h=1/N;x=np.arange(1,N)*h;p=(1+x)**2
   f=(1-p)*np.exp(x)-np.pi**2*np.cos(np.pi*x)+np.pi*p*np.sin(np.pi*x)-1-np.exp(-x)*np.cos(np.pi*x)
   rhs=2*h*h*f;rhs[0]-=(2+h*p[0])*2;rhs[-1]-=(2-h*p[-1])*(np.e-1)
   check(tag+' PDE rows and boundary corrections',np.allclose(b,rhs,atol=1e-15,rtol=0) and np.allclose(A.diagonal(),-4-2*h*h*np.exp(-x)))
  for method,m in case['methods'].items():
   x=np.loadtxt(ROOT/'02_实验本体/results'/(tag+'_'+method+'_solution.csv'),delimiter=',',skiprows=1);res=np.linalg.norm(b-A@x)
   check(tag+' '+method+' saved solution metrics',abs(res-m['residual'])<=max(1e-12,1e-12*res) and abs(np.linalg.norm(x-truth)-m['error_l2'])<=max(1e-12,1e-12*m['error_l2']))
   if m['status']=='converged':check(tag+' '+method+' independent direct reference',np.linalg.norm(x-direct)<1e-6)
   else:check(tag+' '+method+' true divergence',res/np.linalg.norm(b)>1e12)
   if 'grid_n' in case:
    err=x-truth;check(tag+' '+method+' recomputed grid errors',np.isclose(np.max(abs(err)),m['bvp_error_max']) and np.isclose(np.sqrt(np.dot(err,err)/case['grid_n']),m['bvp_error_discrete_l2']))
if ROOT.name.startswith('03_实验03'):
 from fractions import Fraction as F
 R=json.loads((ROOT/'02_实验本体/results/results.json').read_text());q=R['cases']['1.2'];x=np.array(q['x']);y=np.array(q['y']);A=np.column_stack([np.ones(5),x]);c=np.array(q['coefficients'])
 check('Report line normal system exactly matches data',np.array_equal(A.T@A,[[5,151],[151,4629]]) and np.array_equal(A.T@y,[720,22240]))
 check('Report rational coefficients and RSS correct',np.allclose(c,[-3170/43,310/43],atol=1e-12,rtol=0) and abs(q['rss']-4050/43)<1e-11)
 q=R['cases']['1.1_literal'];check('Literal rational system answer correct',np.allclose(q['coefficients'],[1354/890,203/890,-449/890],atol=1e-13,rtol=0) and abs(q['rss']-3721/890)<1e-12)
result=dict(passed=sum(c['passed'] for c in checks),total=len(checks),details=details,checks=checks)
(ROOT/'06_验证记录/delivery_verification.json').write_text(json.dumps(result,indent=2))
print(ROOT.name,result['passed'],'/',result['total'],details)
for c in checks:
 if not c['passed']:print('FAIL',c['name'],c['detail'])
assert result['passed']==result['total']
