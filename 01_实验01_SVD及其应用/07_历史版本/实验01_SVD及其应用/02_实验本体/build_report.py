from pathlib import Path
import json,csv
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parent;O=R/'results';DEST=R.parent/'03_最终实验报告'
def load(n):return json.loads((O/n).read_text())
m=load('matrix_summary.json');im=load('image_summary.json');f=load('eigenfaces_summary.json');q=load('edge_quantization_summary.json');t=load('temporal_summary.json');d=load('synthetic_3dmm_summary.json');run=load('run_summary.json')
doc=Document();sec=doc.sections[0];sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for name in ['Normal','Body Text']:
 st=doc.styles[name];st.font.name='Songti SC';st.font.size=Pt(10);st.element.rPr.rFonts.set(qn('w:eastAsia'),'Songti SC');st.paragraph_format.line_spacing=1.12;st.paragraph_format.space_after=Pt(5)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12)]:
 st=doc.styles[name];st.font.name='Hiragino Sans GB';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('183B56');st.element.rPr.rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
head=sec.header.paragraphs[0];head.text='科学计算  /  实验 01                                      SVD 及其应用';head.style='Caption'
foot=sec.footer.paragraphs[0];foot.alignment=2;foot.add_run('科学计算实验报告  ·  ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
md=[]
def p(s):doc.add_paragraph(s);md.append(s+'\n')
def h(s,level=1):doc.add_heading(s,level);md.append('#'*(level+1)+' '+s+'\n')
def page(s):
 h(s)
 doc.paragraphs[-1].paragraph_format.page_break_before=True
def tab(headers,rows,size=9):
 table=doc.add_table(rows=1,cols=len(headers));table.style='Light Shading Accent 1'
 for c,x in zip(table.rows[0].cells,headers):c.text=str(x)
 for row in rows:
  cells=table.add_row().cells
  for c,x in zip(cells,row):c.text=str(x)
 for row in table.rows:
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  for c in row.cells:
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.space_before=Pt(2)
    for rr in pp.runs:rr.font.size=Pt(size)
 md.append('| '+' | '.join(headers)+' |\n|'+'|'.join(['---']*len(headers))+'|\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
def pic(name,caption,width=6.1):
 doc.add_picture(str(O/name),width=Inches(width));pp=doc.add_paragraph(caption,style='Caption');pp.alignment=1
 md.append(f'![{caption}](../02_实验本体/results/{name})\n')
def code(s):
 pp=doc.add_paragraph();rr=pp.add_run(s);rr.font.name='Menlo';rr.font.size=Pt(8.5)
 md.append('```python\n'+s+'\n```\n')
def fmt(v):return f'{v:.6f}'

doc.add_heading('科学计算实验报告',0);doc.add_heading('第一节上机实验：SVD 及其应用',1)
p('实验日期：2026 年 9 月 17 日\n姓名：____________    学号：____________    班级：____________')
h('摘要',2)
p(f'本实验按给定矩阵完成奇异值分解、截断近似与能量分析，并以 512×512 RGB 宇航员图像完成灰度重建、三通道压缩、全秩范围指标计算和盈亏平衡点验证。在 AT&T（ORL）400 张公开人脸图像上，以训练集学习共享 PCA 基，在独立保留的同身份测试图像上完成最近邻分类、重建及 int8 系数量化。另用合成曲面和模拟系数序列演示 3DMM 与时序共享基思想。全部结果由本地 Python 程序真实运行生成。')
tab(['项目','主要结果'],[
 ['给定矩阵','秩为 4；奇异值约为 100、10、0.1、0.001'],
 ['图像 k=20','PSNR 22.061 dB；参数压缩比 12.788×'],
 ['图像 k=80','PSNR 30.748 dB；参数压缩比 3.197×'],
 ['图像盈亏平衡','k*=255.750；从 k=256 起参数量大于原图'],
 ['Eigenfaces k=80','118/120 正确，准确率 98.33%（固定划分）'],
 ['边缘上传 k=32','每帧 32 B；共享模型原始负载 1,360,384 B'],
 ['执行核验',f'{run["checks_passed"]}/{run["checks_total"]} 项数值和数据检查通过']])
h('完成范围与证据边界',2)
p('核心部分（一、二）全部完成；提高部分完成真实 ORL 数据上的实战 1、实战 3，以及实战 2 的教学替代演示、实战 4 的模拟验证。没有获得真实 3D 扫描或连续视频，因此不把合成曲面称为真实 3D 人脸模型，也不把模拟时序结果称为视频实测。附件没有给出报告模板或学生身份，身份栏留空。')
p('提交目录按“01_指导资料 / 02_实验本体 / 03_最终实验报告”分类。报告可编辑 DOCX 与 PDF 内容一致；完整指标、模型、代码及复现笔记本在实验本体中。')

page('1  实验要求、数据与方法')
h('1.1  要求对应关系',2)
tab(['指导要求','实现与结果文件'],[['一：SVD 与秩一展开','run_experiment.py；matrix_decomposition.npz；matrix_energy.csv'],['Task 1：灰度 k=20','01_grayscale_rank20.png'],['Task 2：RGB 网格','02_rgb_rank_grid.png（原图 + 11 个 k）'],['Task 3：全部 k 指标','image_metrics_all_k.csv（512 行）'],['Task 4：三条曲线','03 / 04 / 05 三张曲线图'],['Task 5：三个发现','报告第 5 节；06_detail_crop.png'],['实战 1 / 3','face_split.csv；eigenfaces_metrics.csv；边缘编码文件'],['实战 2 / 4','synthetic_3dmm_metrics.csv；temporal_summary.json']])
h('1.2  数据来源与预处理',2)
p('图像压缩使用 scikit-image 官方仓库 v0.25.2 的 astronaut.png，尺寸 512×512×3；该图源自 NASA，官方文档标注为公共领域图像 [3]。读取后转换为 float64 并除以 255，使像素处于 [0,1]。灰度取 0.2126R+0.7152G+0.0722B，不缩放原图。')
p('人脸数据下载自剑桥大学托管的 AT&T 官方归档，共 40 人、每人 10 张、92×112 像素、PGM 格式；感谢 AT&T Laboratories Cambridge [2]。直接使用官方原图，没有额外旋转对齐、数据增强或逐图亮度归一化。完整下载 URL、文件大小和 SHA-256 已保存到 data/sources.json。')
h('1.3  SVD 与最优截断',2)
p('对 A∈R^(m×n)，约化分解为 A=UΣVᵀ，σ₁≥…≥σᵣ≥0。U 的列是 AAᵀ 的特征向量，V 的列是 AᵀA 的特征向量，对应特征值为 σᵢ²。非方阵 A 本身没有通常意义的左右“特征向量”，本报告使用“左右奇异向量”。')
p('A=Σᵢ σᵢuᵢvᵢᵀ；Aₖ=Σᵢ₌₁ᵏ σᵢuᵢvᵢᵀ。由正交性，||A−Aₖ||²_F=Σᵢ₍ᵢ₊ₖ₎ σᵢ²，其中求和实际指 i>k。Aₖ 是 Frobenius 范数意义下的最优秩不超过 k 的近似。截断是数据相关的低秩近似，不等同于按傅里叶频率排序。')

page('2  给定矩阵的分解与数值核对')
h('2.1  组装约定',2)
p(f'按 PDF 图像逐行转录 Q₅、Q₄、Σ 后，比较两种乘法。Q₅ΣQ₄ 与指导书印出的 A 的最大元素差为 {m["assembled_vs_printed_max_abs"]:.3e}，而 Q₅ΣQ₄ᵀ 的差为 {m["transpose_candidate_vs_printed_max_abs"]:.6f}。因此采用 A=Q₅ΣQ₄；此处 Q₄ 对应 Vᵀ。两者的小差异与指导矩阵只保留六位小数相符，未将输入擅自重新正交化。')
p(f'指导书将 Q 称为“正定矩阵”，但它们非对称、实际近似正交。||Q₅ᵀQ₅−I||_F={m["Q5_orthogonality_error"]:.3e}，||Q₄ᵀQ₄−I||_F={m["Q4_orthogonality_error"]:.3e}。精确输入与秩一分量保存在 matrix_decomposition.npz。')
tab(['A 列1','列2','列3','列4'],[[f'{x:.6f}' for x in row] for row in m['A']])
h('2.2  秩、奇异值与能量',2)
tab(['i / k','σᵢ','E(k)','相对误差'],[[r['k'],f'{m["singular_values"][i]:.9f}',f'{r["energy"]:.12f}',f'{r["relative_error"]:.6e}'] for i,r in enumerate(m['energy'])])
p(f'4 个奇异值均为正，所以秩为 4（满列秩）。κ₂(A)≈{m["condition_number"]:.3f}，跨度达五个数量级。保留前 2 项已包含约 99.999901% 的能量，但原矩阵仍严格为秩 4，不能把“能量接近 100%”当成精确低秩。AAᵀ 还有一个零特征值，AᵀA 则无零特征值。')
h('2.3  两种算法交叉验证',2)
p(f'教学实现先用 eigh 求 AᵀA 的特征对，取平方根得到 σᵢ，再按 uᵢ=Avᵢ/σᵢ 计算左奇异向量；参考实现使用 np.linalg.svd。两者奇异值最大绝对差 {m["eigen_implementation_max_s_error"]:.3e}。秩一分量求和重构误差为 {m["reconstruction_error"]:.3e}。')
p('构造 AᵀA 会将条件数平方，因此该实现适合解释数学关系；小奇异值和左向量正交性更易受舍入影响，后续图像与人脸实验均直接用 SVD。奇异向量允许成对变号，同一 σᵢ 对应的 uᵢ、vᵢ 同时取负不会改变重构。附录给出本次输出的向量与全部秩一项。')

page('3  图像压缩：实现与指标')
h('3.1  灰度图三行核心代码',2)
code('A = img @ np.array([0.2126, 0.7152, 0.0722])\nU, s, Vt = np.linalg.svd(A, full_matrices=False)\nA_k = (U[:, :20] * s[:20]) @ Vt[:20, :]')
p('输入 img 已转换为 [0,1] 浮点数。输出 A_k.shape=(512,512)。显示时使用 np.clip(A_k,0,1)，不以裁剪后矩阵计算奇异值尾和。full_matrices=False 返回约化形状；它并不意味着只计算前 k 个奇异值 [4]。')
pic('01_grayscale_rank20.png','图 1  灰度原图与 k=20 重建',4.5)
h('3.2  三通道合计公式',2)
p('R、G、B 独立分解。设 T=Σ_cΣ_i σ²_(c,i)，Rₖ=Σ_cΣ_(i>k)σ²_(c,i)，则 E(k)=1−Rₖ/T，relerr=√(Rₖ/T)，MSE=Rₖ/(HWC)，PSNR=10log₁₀(1/MSE)。每通道需 Hk+k+kW 个标量，参数压缩比 CR=HW/[k(H+W+1)]。')
p('stats(k) 直接读取奇异值平方的反向累积尾和，不重建、不逐像素相减。逆序累加避免用总能量减前缀能量产生消减误差。满秩时定义尾和为 0、理论 PSNR=∞；显式浮点重构仍可能留下约机器精度的残差。')
tab(['k','能量 (%)','相对误差','PSNR / dB','参数压缩比'],[[r['k'],f'{r["energy"]*100:.4f}',f'{r["relative_error"]:.6f}','∞' if isinstance(r['psnr_db'],str) else f'{r["psnr_db"]:.3f}',f'{r["compression_ratio"]:.3f}'] for r in im['selected']])

page('4  RGB 重建与三条曲线')
pic('02_rgb_rank_grid.png','图 2  原图与 k∈{1,2,3,5,8,12,20,35,50,80,120} 的 RGB 重建',6.35)
p('k=1～4 时，整体亮暗与大色块仍存在，但五官、衣服徽章及器件纹理不可辨认；k=20 后轮廓较清楚，仍有明显横纵纹理失真；k=80 与 120 时主要细节逐步恢复。图像按 [0,1] 裁剪后显示，标题 PSNR 来自裁剪前理论误差。')
p('所有 k=1,…,512 的完整指标均在 image_metrics_all_k.csv。报告展示的是摘要行，未省略实验输出。另以 1、4、20、80、120、256、512 七个 k 显式重建，验证尾和 MSE 与逐像素误差一致；裁剪后的 MSE 另存，避免混用口径。')

page('4.1  奇异值、质量与存储代价曲线')
pic('03_singular_values.png','图 3  三通道奇异值的对数衰减曲线',4.7)
pic('04_error_psnr.png','图 4  相对误差下降，PSNR 上升；满秩 ∞ 不绘为有限点',4.7)
pic('05_compression_ratio.png','图 5  参数压缩比与盈亏平衡点',4.7)

page('5  三个关键发现及解释')
h('5.1  能量陷阱',2)
p('本图 k=1 的能量占比为 78.5057%，k=4 为 90.5566%，但 PSNR 仅 11.856 dB 和 15.428 dB，结构细节已经严重丢失。k=20 能量达 97.9497%，PSNR 仍只有 22.061 dB。能量是平方和尺度，大面积背景和平均亮度贡献显著；不能仅凭能量比例判断视觉质量。')
h('5.2  盈亏平衡点',2)
p('CR>1 ⇔ HW>k(H+W+1) ⇔ k<k*=HW/(H+W+1)。本图 k*=262144/1025=255.750244，故 k≤255 时参数更少，k≥256 时参数更多。指导书的约 353 对应另一个 620×820 尺寸，本实验不沿用该数值。')
p('上述比较假设每个标量占相同字节。原图 uint8 原始负载 786,432 B；k=20 的因子若按 float32 保存，需 246,000 B，字节压缩比为 3.197×，而不是 12.788×。若不再量化，float32 因子的字节盈亏平衡点约为 k=63.938，从 k=64 就开始大于原始 uint8 像素负载。原 PNG 文件为 791,555 B，其大小还受到编码与元数据影响。这里未声称 SVD 优于 JPEG。')
h('5.3  细节损失',2)
pic('06_detail_crop.png','图 6  同一区域局部对照：原图、k=1、k=4、k=20',6.1)
p('宇航服徽章、衣服接缝和器件边缘比大块橙色区域更早变得模糊，说明该样本的细节需要更多奇异分量共同表达。小奇异值并不普遍等价于高空间频率；“高频细节先损失”是本图观察到的趋势，不是每张图的严格定理。')

page('6  Eigenfaces：训练与测试设计')
p('Eigenfaces 将一张脸展开为 D=112×92=10,304 维向量。按 Turk 与 Pentland 的共享子空间思路 [5]，对训练矩阵求均值 μ，中心化后 X_c=UΣB；B 的前 k 行作为共同基。样本系数 z=(x−μ)Bₖᵀ，重建 x̂=μ+zBₖ，分类在系数空间中使用欧氏距离最近邻。')
p('随机种子固定为 20260917，每个人独立随机选 7 张训练、3 张测试，共 280/120 张。均值、PCA 基、量化参数只由训练集得到；测试图像不参与拟合。所有划分保存在 face_split.csv。这是已知身份的闭集识别，测试人员与训练人员相同，测试照片不同，不能解释为对未见过身份的识别能力。')
p('主评估 k=80 在运行前固定；其他 k 用于敏感性分析，不以测试表现选择新的“最佳参数”。单次划分结果具有样本依赖性，不复述指导书的约 96% 为本次结果。')
pic('08_mean_eigenfaces.png','图 7  训练集均值脸及前 11 个特征脸；基图分别归一化显示',4.6)
p('中心化训练矩阵数值秩为 279；特征脸正交性残差为 '+f'{f["orthogonality_error"]:.3e}'+'。特征脸是带正负值的空间模式，不是另外采集的真实人脸照片。')

page('6.1  Eigenfaces 的识别和重建结果')
tab(['k','正确 / 120','准确率 (%)','测试 PSNR','训练库参数压缩比'],[[r['k'],r['correct'],f'{r["accuracy"]*100:.2f}',f'{r["test_psnr_db"]:.3f}',f'{r["train_library_parameter_ratio"]:.3f}'] for r in f['metrics']])
p('k=80 的主结果为 118/120，准确率 98.33%。样本级 Wilson 95% 区间约 [94.13%,99.54%]，仅作描述；同一身份内样本相关，不能把这个区间当作跨身份泛化置信保证。逐张预测已保存，错误样本可追溯。')
pic('09_face_reconstructions.png','图 8  三张保留测试图像在不同 k 下的重建',4.9)
p('k=200、279 的重建 PSNR 继续提高，识别准确率却下降至 96.67%，说明保真与判别不是同一个优化目标。测试数据在训练子空间外仍有分量，因此即使使用训练集全部 279 个非零方向，测试重建也不是无误差。')
p('含均值与共享基时，N 张图的参数压缩比为 ND/[D+kD+Nk]。本训练库 N=280、k=80 的比值是 3.366×。公式中的均值和基不能忽略；若基为 float32、原图为 uint8，实际字节比较还需计入 4 倍字长。')

page('7  边缘设备共享基与 int8 量化')
p('采用 k=32 的同一训练 PCA 基，按每个系数维度的训练分布 0.1% / 99.9% 分位数定义 lo、hi，并令 scale=(hi−lo)/255。编码 q=round(clip((z−lo)/scale,0,255))−128，q 为 int8；解码 ẑ=(q+128)scale+lo，再重建 μ+ẑB。没有使用测试分布标定量化器。')
tab(['项目','本次实测'],[['浮点系数重建 PSNR',f'{q["float_psnr_db"]:.6f} dB'],['int8 系数重建 PSNR',f'{q["int8_psnr_db"]:.6f} dB'],['PSNR 损失',f'{q["psnr_loss_db"]:.6f} dB'],['相对浮点重建的最大像素差',f'{q["max_pixel_difference"]:.6f}（像素范围 0～1）'],['测试系数超出分位范围',f'{q["clipped_coefficient_fraction"]*100:.4f}%'],['量化后最近邻准确率',f'{q["int8_recognition_accuracy"]*100:.2f}%（训练和测试均量化）'],['原始 uint8 / 系数帧负载','10,304 B / 32 B（322×）'],['共享模型原始 / NPZ 文件','1,360,384 B / 1,362,124 B'],['含共享模型的 120 帧压缩比',f'{q["first_120_frames_ratio"]:.4f}×'],['相对原始像素的摊销盈亏平衡',f'{q["break_even_frames"]} 帧']])
p('平均 PSNR 变化很小，但最大像素差达到 0.0211，明显不等于指导书示例的 0.0015。重尾样本的分位截断可能产生局部较大误差，不能仅用平均指标宣称量化没有代价。')
p('模型原始负载包含 32×10304 个基元素、10304 个均值、32 个 lo 和 32 个 scale，均按 float32 计。总发送量为 B_model+32N；相对原始 uint8 帧节省要求 N> B_model/(10304−32)，本次至少 133 帧。该估算未计网络包头、身份元数据、加密与重传开销。')
p('已输出实际 int8 编码、共享基与模型版本，并对模型原始字节计算 SHA-256 标识；本地加载已保存的 float32 模型解码后，与计算结果交叉核验。部署时需比对基 ID、版本和量化参数；本实验完成的是本地编解码，没有声称已部署到真实边缘设备。')
p('系数化不是加密，也不能保证去标识化：本实验恰恰能用系数进行身份分类。持有基的一方能够近似恢复人脸；没有基也不构成可证明的隐私保护。去掉若干主分量的身份保护效果并未在本实验中验证。')

page('8  3DMM 教学替代与时序模拟')
h('8.1  共享基如何推广到三维',2)
p('3DMM 将对应顶点的形状坐标和纹理分别展开，使用 S=μ_S+Σᵢαᵢbᵢ、T=μ_T+Σᵢβᵢcᵢ 表示 [6]。本次未获取受限真实 3D 扫描，按照数据说明允许的教学替代路径，以 Eigenfaces 验证共同数学机制，并额外构造合成曲面检验形状系数化。')
p('合成曲面为 40×32 网格，1,280 个顶点、3,840 维坐标。用宽度、高度、局部隆起等 6 个线性形变模式和固定种子生成 200 个样本，160 个训练、40 个测试。该曲面不代表真实人脸，不具备真实 3DMM 的解剖或逼真度意义，也未实现从照片反演 3D 或纹理 PCA。')
pic('11_synthetic_3dmm.png','图 9  合成曲面及低秩重建（明确不是受试者 3D 扫描）',6.1)
tab(['k','保留测试曲面 RMSE'],[[r['k'],f'{r["test_rmse"]:.6e}'] for r in d['metrics']])
p('k=6 时恢复到浮点精度，因为数据的生成机制本来只有 6 个线性自由度。这是程序正确性演示，不能作为真实三维人脸只需要 6 个参数的证据。')
h('8.2  时序模拟条件',2)
p('ORL 是静态照片库，未提供连续采集的时间戳，不能把文件编号当视频顺序。因此以一个测试样本的 PCA 系数为中心，叠加振幅为各训练系数标准差 0.12 倍的平滑正弦变化，生成 300 帧、32 维模拟轨迹。量化仍使用训练参数。此构造主动满足平滑性，结论仅适用于该模拟条件。')

page('8.3  差分编码的实际字节核验')
pic('10_synthetic_temporal_coefficients.png','图 10  模拟系数轨迹的前四维',4.2)
p('每 60 帧保存一次绝对系数，其余帧记录相邻差；差分可能超出 int8 范围，故按 int16 保存，再使用 zlib level=9 压缩。逐帧累计恢复后的量化系数与原系数逐元素完全一致，zlib 解压字节亦一致。这里验证的是无损恢复量化系数，不是对原图无损。')
tab(['数据表示','字节数'],[['直接 int8 系数',t['direct_int8_bytes']],['未压缩 int16 差分',t['delta_int16_bytes']],['直接 int8 + zlib',t['direct_zlib_bytes']],['int16 差分 + zlib',t['delta_zlib_bytes']]])
p(f'公平比较双方均采用 zlib：直接系数 1,979 B，差分系数 1,387 B，额外压缩为 {t["extra_factor_vs_direct_zlib"]:.3f}×，没有达到指导书举例的 2～3×。如果只做差分而不熵编码，int16 反而把负载扩大一倍。包含一次共享模型分发时，300 帧相对原始像素负载的压缩比为 {t["ratio_with_model_vs_raw_pixels"]:.3f}×；未计网络封装。')
p('绝对系数刷新帧只重置预测状态，不会重算 PCA 基。本实验没有实测场景切换、姿态漂移检测或真实视频上的基更新，也没有与 JPEG/H.264 等编码器进行率失真比较，因此不能据此宣称这是唯一或普遍更好的视频压缩方式。')
h('9  实验结论',2)
p('（1）给定矩阵的秩、奇异值及能量结构均与设定一致；核对乘法约定是复现实验的必要步骤。\n（2）低秩压缩的质量、能量与存储是不同指标；本图 k=80 达到约 30.75 dB，但按 float32 字节存储时已不比 uint8 原始负载小。\n（3）共享基把单张图像表示为少量系数，同时可用于识别；基与均值的开销必须由足够多样本摊销。\n（4）量化与差分编码有实测收益，也有截断误差和字节开销；真实数据结论与模拟机制验证应分别表述。')

page('10  智能体协作过程与复现')
p('本课程使用智能体完成实验。本次流程为：读取并逐页核对附件 → 提取要求与矩阵 → 获取公开数据及论文 → 编写教学实现和参考实现 → 真实执行 → 核查反例与指标口径 → 自动生成报告 → 核验文档。报告没有虚构课堂分组讨论或人工操作记录。')
p('智能体发现并处理的关键问题包括：PDF 矩阵文本按列抽取导致的转置风险；Q₄ 在给定 A 中是 Vᵀ；近似正交与正定的术语差异；裁剪图像后的误差不再等于原始奇异值尾和；参数量与实际字节量不能混用；静态照片不得伪装为真实时序数据。')
h('10.1  本地执行环境',2)
tab(['条目','版本或设置'],[['Python',run['python']],['NumPy / SciPy',run['numpy']+' / '+run['scipy']],['Matplotlib / Pillow',run['matplotlib']+' / '+run['pillow']],['随机种子',run['seed']],['本次主体计算耗时',f'{run["elapsed_seconds"]:.3f} s（不含下载、依赖安装、字体初始化和报告生成）'],['数值检查',f'{run["checks_passed"]}/{run["checks_total"]} 通过']])
h('10.2  一键复现',2)
code('cd /Users/juleyau/code/sci_comp\n.venv/bin/python 实验01_SVD及其应用/02_实验本体/run_experiment.py\n.venv/bin/python 实验01_SVD及其应用/02_实验本体/build_report.py')
p('数据已随目录保存，重新计算无需联网。换电脑时先创建 Python 虚拟环境并安装 requirements.txt；requirements-lock.txt 保存本次确切环境。experiment.ipynb 提供按模块运行与结果查看入口。run.log 是真实运行记录，verification.json 列出自动检查，manifest_sha256.csv 用于交付文件完整性核对。')
h('10.3  尚不能由本次实验回答的问题',2)
p('真实 3D 人脸的拟合精度、真实视频的码率收益、不同随机划分和新身份下的识别泛化，以及隐私攻击下的保护程度，均未验证。报告仅将已运行的核心实验、ORL 识别与量化实验，以及明确标注的教学模拟作为结果。')

page('参考资料与数据出处')
refs=[
'[1] 课程附件：《第一节上机实验——SVD 及其应用》《数据集获取说明》。本次要求以这两份附件为准；原件保存在 01_指导资料。',
'[2] AT&T Laboratories Cambridge. The Database of Faces. https://www.cl.cam.ac.uk/research/dtg/attarchive/facedatabase.html （访问：2026-09-17）。数据署名：AT&T Laboratories Cambridge。',
'[3] scikit-image. skimage.data.astronaut. https://scikit-image.org/docs/stable/api/skimage.data.html#skimage.data.astronaut （访问：2026-09-17）。本次输入固定为官方仓库 v0.25.2 的 astronaut.png，原始图源 NASA。',
'[4] NumPy. numpy.linalg.svd. https://numpy.org/doc/stable/reference/generated/numpy.linalg.svd.html （访问：2026-09-17）。',
'[5] Turk M, Pentland A. Eigenfaces for Recognition. Journal of Cognitive Neuroscience, 1991, 3(1): 71–86. DOI: 10.1162/jocn.1991.3.1.71. MIT 课程存档：https://courses.media.mit.edu/2008fall/mas622j/Turk%2BPentland-1991.pdf 。',
'[6] Blanz V, Vetter T. A Morphable Model for the Synthesis of 3D Faces. Proceedings of SIGGRAPH 1999: 187–194. DOI: 10.1145/311535.311556. 作者项目归档：https://shapemodelling.cs.unibas.ch/gravis-site-archive/Sigg99.html 。论文副本已保存。',
'[7] OpenCV. Face Recognition with OpenCV. https://docs.opencv.org/4.13.0/da/d60/tutorial_face_main.html （访问：2026-09-17）。本实验自行实现 PCA 最近邻流程，未直接调用 OpenCV 识别器。',
'[8] University of Basel. Basel Face Model. https://faces.dmi.unibas.ch/bfm/bfm2019.html （访问：2026-09-17）。仅用于真实三维模型背景，本次未下载 BFM 数据。']
for r in refs:p(r)
p('说明：文献用于方法背景和数据来源，报告中所有“本次”数值均来自 results 中的运行产物。没有引用动态论文被引次数，也没有将课程材料中的其他实验数字当成本次测量。')

page('附录 A  左右奇异向量')
p('下面是组装 A 的约化分解结果，保留 6 位小数。U 的第 i 列为左奇异向量 uᵢ；Vᵀ 的第 i 行转置后为右奇异向量 vᵢ。因保留位数有限，用此表重新相乘可能产生微小舍入误差；精确数组在 NPZ 文件中。')
h('A.1  U（5×4）',2);tab(['u₁','u₂','u₃','u₄'],[[f'{x:+.6f}' for x in row] for row in m['U']],10)
h('A.2  Vᵀ（4×4）',2);tab(['第1列','第2列','第3列','第4列'],[[f'{x:+.6f}' for x in row] for row in m['Vt']],10)
p('对应关系为 Avᵢ=σᵢuᵢ，Aᵀuᵢ=σᵢvᵢ；AᵀAvᵢ=σᵢ²vᵢ，AAᵀuᵢ=σᵢ²uᵢ。σᵢ² 近似为 10000、100、0.01、0.000001。完整 5×5 的 AAᵀ 另有一个零特征值方向，不属于本次约化 U 的四列。')
p('指导书印出的 A 自身也是六位小数输入。对其直接 SVD 的奇异值为：'+', '.join(f'{v:.9f}' for v in m['printed_A_singular_values'])+'。它与从 Q、Σ 组装的版本存在轻微差异，因此结果文件同时保留 A 与 A_printed，主报告统一分析组装版本。')

page('附录 B  秩一矩阵之和')
import numpy as np
terms=np.load(O/'matrix_decomposition.npz')['rank_one_terms']
p('记 Rᵢ=σᵢuᵢvᵢᵀ，则 A=R₁+R₂+R₃+R₄。下列四项均为 5×4 秩一矩阵；Aₖ 即前 k 项之和。')
for i,term in enumerate(terms):
 h(f'R{i+1} = σ{i+1} u{i+1} v{i+1}ᵀ',2)
 tab(['第1列','第2列','第3列','第4列'],[[f'{x:+.7f}' for x in row] for row in term],8)
# 修正排版中的数学尾和为直观、无歧义的纯文本。
for para in doc.paragraphs:
 if 'Σᵢ₍ᵢ₊ₖ₎' in para.text:
  para.text=para.text.replace('Σᵢ₍ᵢ₊ₖ₎ σᵢ²，其中求和实际指 i>k','Σ_(i>k) σᵢ²')
md=[s.replace('Σᵢ₍ᵢ₊ₖ₎ σᵢ²，其中求和实际指 i>k','Σ_(i>k) σᵢ²') for s in md]
DEST.mkdir(exist_ok=True);doc.save(DEST/'实验01_SVD及其应用_实验报告.docx');(DEST/'实验01_SVD及其应用_实验报告.md').write_text('\n'.join(md),encoding='utf-8')
print('Report saved',DEST)
