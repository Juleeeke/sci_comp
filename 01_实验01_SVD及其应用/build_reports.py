#!/usr/bin/env python3
"""Create two English reports from the same measured results; retain the supplied template."""
from pathlib import Path
import json, shutil, re, sys
from collections import Counter
from lxml import etree
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT/'05_构建工具'))
math_cache_path=ROOT/'03_最终实验报告/equations.json'
stored_math=json.loads(math_cache_path.read_text()) if math_cache_path.exists() else {}
math_cache=stored_math.get('equations',{}) if stored_math.get('version')==2 else {}
def office_math(latex):
    if latex not in math_cache:
        from latex2mathml.converter import convert
        transform=etree.XSLT(etree.parse('/Applications/Microsoft Word.app/Contents/Resources/mathml2omml.xsl'))
        ml=etree.fromstring(convert(latex).encode())
        ns='{http://www.w3.org/1998/Math/MathML}'
        # The Office stylesheet expects grouped delimiter contents in mfenced.
        # Modern converter mrows can otherwise lose tokens such as A-A_k or clip.
        pairs={'(':')','[':']','{':'}','‖':'‖','|':'|'}
        for row in list(ml.iter(ns+'mrow'))[::-1]:
            if len(row)>=2 and row[0].tag==ns+'mo' and row[-1].tag==ns+'mo' and pairs.get(row[0].text)==row[-1].text:
                if row[0].get('stretchy')=='true' or row[0].get('fence')=='true':
                    fenced=etree.Element(ns+'mfenced',open=row[0].text,close=row[-1].text,separators='')
                    inner=etree.SubElement(fenced,ns+'mrow')
                    for child in list(row)[1:-1]:inner.append(child)
                    row.getparent().replace(row,fenced)
        result=transform(ml)
        tokens=lambda s:Counter(re.findall(r'[A-Za-z0-9\u0370-\u03ff−]',s))
        before=tokens(''.join(ml.itertext()))
        after=tokens(''.join(x.text or '' for x in result.getroot().iter(qn('m:t'))))
        if before!=after:raise ValueError('Math conversion lost or added a semantic token: '+latex)
        math_cache[latex]=etree.tostring(result.getroot(),encoding='unicode')
    return etree.fromstring(math_cache[latex].encode())
OUT = ROOT/'02_实验本体/results'
def load(n): return json.loads((OUT/n).read_text())
m=load('matrix_summary.json'); im=load('image_summary.json'); f=load('eigenfaces_summary.json')
q=load('edge_quantization_summary.json'); t=load('temporal_summary.json'); d=load('synthetic_3dmm_summary.json'); run=load('run_summary.json')
npz=np.load(OUT/'matrix_decomposition.npz')
pages=[]
def page(title):
    pages.append({'title':title,'blocks':[]})
def add(kind, **kw): pages[-1]['blocks'].append(dict(kind=kind,**kw))
def p(s): add('paragraph',text=s)
def h(s): add('heading',text=s)
def eq(word,latex): add('equation',word=word,latex=latex)
def tab(headers,rows,size=9): add('table',headers=headers,rows=[[str(x) for x in r] for r in rows],size=size)
def fig(name,caption,height): add('figure',name=name,caption=caption,height=height)
def code(s): add('code',text=s)
def mat(title,a,precision=6):
    h(title);tab([f'Column {j+1}' for j in range(a.shape[1])],[[f'{x:+.{precision}f}' for x in r] for r in a],8.5)

page('1. Objectives, data and matrix construction')
p('The experiment investigates singular value decomposition (SVD), optimal low-rank approximation, image compression, and shared-basis representations. All five image tasks in the course handout are implemented: grayscale reconstruction, the RGB rank grid, metrics for every rank, three diagnostic curves, and an evidence-based discussion of the three key findings. The extension covers Eigenfaces, coefficient quantization, a synthetic three-dimensional surface, and synthetic temporal coefficients.')
p('The RGB input is the 512 x 512 astronaut image distributed with scikit-image v0.25.2 (NASA image credit). Pixel values are converted to float64 and divided by 255. Eigenfaces uses the archived AT&T/ORL database: 400 real grayscale photographs, 40 subjects, and 112 x 92 pixels per image [2, 3]. The original inputs are included with the experiment.')
h('1.1 Assembly and consistency of the supplied matrices')
p(f'The numerical matrix printed in the handout agrees with A = Q5 Sigma Q4, where Q4 acts as V-transpose. The maximum elementwise difference is {m["assembled_vs_printed_max_abs"]:.3e}; replacing Q4 by its transpose gives a difference of {m["transpose_candidate_vs_printed_max_abs"]:.6f}. The first discrepancy is consistent with six-decimal input rounding. The main experiment uses the assembled matrix; the printed matrix is retained separately for comparison. Exact input arrays appear in the saved NPZ and Appendix A.')
mat('Assembled A (5 x 4)',np.array(m['A']))
p(f'The handout calls Q5 and Q4 positive definite, but these nonsymmetric matrices are approximately orthogonal: their Frobenius orthogonality residuals are {m["Q5_orthogonality_error"]:.3e} and {m["Q4_orthogonality_error"]:.3e}. No reorthogonalization is applied to the given rounded entries.')
h('1.2 Mathematical representation')
eq('A = UΣVᵀ = Σᵢ σᵢuᵢvᵢᵀ;   Aₖ = Σᵢ₌₁ᵏ σᵢuᵢvᵢᵀ',r'A=U\Sigma V^T=\sum_{i=1}^{r}\sigma_i u_i v_i^T,\qquad A_k=\sum_{i=1}^{k}\sigma_i u_i v_i^T.')
p('For a real m x n matrix, reduced SVD returns U of size m x r, V-transpose of size r x n and descending singular values, with r = min(m,n). Retaining the leading k terms gives a best rank-at-most-k approximation in the Frobenius norm. Reduced output is distinct from truncation: full_matrices=False controls returned shapes; slicing selects k [4].')

page('2. Matrix singular values, rank and energy')
p('A teaching implementation obtains the eigenpairs of A-transpose A using numpy.linalg.eigh, orders them by descending eigenvalue, takes square roots to obtain singular values, and forms each left vector as A v / sigma. It is cross-checked against numpy.linalg.svd. Forming A-transpose A squares the condition number, so direct SVD is used for the image and face experiments.')
eq('Avᵢ = σᵢuᵢ;  Aᵀuᵢ = σᵢvᵢ;  AᵀAvᵢ = σᵢ²vᵢ;  AAᵀuᵢ = σᵢ²uᵢ',r'Av_i=\sigma_i u_i,\quad A^Tu_i=\sigma_i v_i,\quad A^TAv_i=\sigma_i^2v_i,\quad AA^Tu_i=\sigma_i^2u_i.')
p('Because A is rectangular, u and v are singular vectors of A. They are eigenvectors of A A-transpose and A-transpose A, respectively; they are not ordinary eigenvectors of the rectangular matrix itself.')
tab(['k','Singular value','Retained energy E(k)','Relative error'],[[r['k'],f'{m["singular_values"][i]:.9f}',f'{r["energy"]:.12f}',f'{r["relative_error"]:.3e}'] for i,r in enumerate(m['energy'])])
eq('E(k) = (Σᵢ≤ₖ σᵢ²)/(Σᵢ σᵢ²);  ||A−Aₖ||F/||A||F = √(1−E(k))',r'E(k)=\frac{\sum_{i\le k}\sigma_i^2}{\sum_i\sigma_i^2},\qquad \frac{\left\|A-A_k\right\|_F}{\left\|A\right\|_F}=\sqrt{1-E(k)}.')
p(f'All four singular values are nonzero, so A has rank 4 and full column rank. Its two-norm condition number is {m["condition_number"]:.3f}. The first two components retain about 99.999901% of the energy, but this does not change the exact rank. A-transpose A has four positive eigenvalues, approximately 10000, 100, 0.01 and 0.000001. A A-transpose has one additional zero eigenvalue.')
p(f'The two implementations differ in singular values by at most {m["eigen_implementation_max_s_error"]:.3e}. Summing all four rank-one terms gives a reconstruction residual of {m["reconstruction_error"]:.3e}. The measured relative reconstruction error at k=4 is a floating-point residual; the mathematical truncated-SVD tail is exactly zero.')
h('Left and right singular vectors')
mat('Reduced U: each column is a left singular vector',np.array(m['U']))
mat('V-transpose: each row transposes to a right singular vector',np.array(m['Vt']))
p('A paired sign change in a left and right vector leaves its rank-one term unchanged. Tables are rounded for readability; full precision is retained in matrix_decomposition.npz.')

page('3. Explicit sum of rank-one matrices')
eq('Rᵢ = σᵢuᵢvᵢᵀ;  A = R₁ + R₂ + R₃ + R₄',r'R_i=\sigma_i u_i v_i^T,\qquad A=R_1+R_2+R_3+R_4.')
p('The four components below are the requested explicit rank-one decomposition. Each component is 5 x 4, and A_k is obtained by adding the first k components. The full-precision arrays are also saved individually as CSV files.')
for i,a in enumerate(npz['rank_one_terms']): mat(f'R{i+1}',a,7)
p('The first component carries almost all of the large-scale magnitude. Later terms are much smaller, yet the fourth term is essential for exact rank 4. A visually negligible or numerically small component is not identical to a zero component.')

page('4. Grayscale reconstruction and metric definitions')
h('4.1 Three-line core implementation (Task 1)')
code('A = img @ np.array([0.2126, 0.7152, 0.0722])\nU, s, Vt = np.linalg.svd(A, full_matrices=False)\nA_k = (U[:, :20] * s[:20]) @ Vt[:20, :]')
p('The input img is an RGB floating-point array in [0,1]. The weighted sum produces a grayscale matrix. For k=20, the actual reconstruction shape is (512,512). The reconstruction is clipped to [0,1] for display only; raw error identities use the unclipped matrix.')
fig('01_grayscale_rank20.png','Figure 1. Original grayscale image and its rank-20 reconstruction.',2.3)
h('4.2 Metrics without explicit reconstruction (Task 3)')
p('Let T be the sum of squared singular values over all three color channels, and let R(k) be the sum of squared singular values discarded after rank k. Orthogonal invariance of the Frobenius norm gives the following formulas, with pixel peak value 1:')
eq('T = ΣcΣᵢ σc,ᵢ²;  R(k) = ΣcΣᵢ>ₖ σc,ᵢ²;  E(k) = 1−R(k)/T',r'T=\sum_c\sum_i\sigma_{c,i}^2,\quad R(k)=\sum_c\sum_{i>k}\sigma_{c,i}^2,\quad E(k)=1-\frac{R(k)}{T}.')
eq('Relative error = √(R(k)/T);  MSE = R(k)/(HWC);  PSNR = 10 log₁₀(1/MSE)',r'\mathrm{relerr}=\sqrt{\frac{R(k)}{T}},\quad \mathrm{MSE}=\frac{R(k)}{HWC},\quad \mathrm{PSNR}=10\log_{10}\frac{1}{\mathrm{MSE}}.')
eq('Parameters = Ck(H+W+1);  CR = HW/[k(H+W+1)]',r'P(k)=Ck(H+W+1),\qquad \mathrm{CR}(k)=\frac{HW}{k(H+W+1)}.')
p('For each channel, U needs Hk scalars, singular values need k, and V-transpose needs kW. The +1 term accounts for the singular values. Reverse cumulative sums compute the tails and avoid cancellation from subtracting nearly equal total and prefix energies. At full rank the theoretical MSE is zero and PSNR is infinity.')

page('5. RGB rank grid and numerical results')
p('Each RGB channel is decomposed separately, truncated at the same k, and stacked to reconstruct a color image (Task 2). The prescribed ranks are 1, 2, 3, 5, 8, 12, 20, 35, 50, 80 and 120. The original plus these eleven reconstructions form the required 3 x 4 grid.')
fig('02_rgb_rank_grid.png','Figure 2. RGB reconstruction grid. Display values are clipped; titles use theoretical, unclipped PSNR.',4.4)
tab(['k','Energy (%)','Rel. error','PSNR (dB)','Parameter CR'],[[r['k'],f'{r["energy"]*100:.4f}',f'{r["relative_error"]:.6f}','Infinity' if isinstance(r['psnr_db'],str) else f'{r["psnr_db"]:.3f}',f'{r["compression_ratio"]:.3f}'] for r in im['selected']],8.5)
p('All 512 rows, for k=1 through 512, are supplied in image_metrics_all_k.csv. Seven independent explicit reconstructions at k=1,4,20,80,120,256,512 confirm the singular-value tail identity against pixelwise error. Clipped MSE is recorded separately in image_error_validation.csv, because clipping changes the error measure.')

page('6. Diagnostic curves and break-even rank (Task 4)')
fig('03_singular_values.png','Figure 3. Singular values of R, G and B on a logarithmic vertical scale.',2.15)
fig('04_error_psnr.png','Figure 4. Relative error decreases and PSNR increases with k. The infinite full-rank PSNR is omitted from the finite curve.',2.15)
fig('05_compression_ratio.png','Figure 5. Parameter compression ratio and the rank at which the parameter count equals the pixel count.',2.15)
p('The singular-value spectra decay rapidly, allowing small ranks to represent broad structures. Quality improves as k increases, while the parameter compression ratio falls. Once k exceeds the break-even rank, the factor representation contains more scalars than the original image.')

page('7. Three key findings (Task 5)')
h('7.1 Energy trap')
selected={r['k']:r for r in im['selected']}
p(f'At k=1, retained energy is {selected[1]["energy"]*100:.4f}% but PSNR is only {selected[1]["psnr_db"]:.3f} dB. At k=4, energy reaches {selected[4]["energy"]*100:.4f}% while PSNR is {selected[4]["psnr_db"]:.3f} dB. Even k=20 retains {selected[20]["energy"]*100:.4f}% with PSNR {selected[20]["psnr_db"]:.3f} dB. Large color regions and global brightness dominate squared energy; high retained energy alone does not imply good visual fidelity.')
h('7.2 Break-even point: scalars versus actual bytes')
eq('CR > 1 iff k < k* = HW/(H+W+1)',r'\mathrm{CR}>1\iff k<k^*=\frac{HW}{H+W+1}.')
p(f'For the actual 512 x 512 image, k* = {im["kstar"]:.6f}. Therefore k<=255 reduces the scalar count and k>=256 increases it. The handout\'s approximately 353 belongs to a different 620 x 820 image and is not reused as a measured value.')
p(f'This comparison assumes equal bytes per scalar. The RGB uint8 pixel payload is {im["raw_uint8_bytes"]:,} B. Rank-20 float32 factors need {im["k20_float32_factor_bytes"]:,} B, giving a byte ratio of {im["raw_uint8_bytes"]/im["k20_float32_factor_bytes"]:.3f}, compared with the parameter ratio {selected[20]["compression_ratio"]:.3f}. The float32-versus-uint8 break-even rank is only {im["kstar"]/4:.3f}; k>=64 expands the raw byte payload. The source PNG is {im["source_png_bytes"]:,} B and has its own encoding overhead. No comparison with JPEG was performed.')
h('7.3 Detail loss')
fig('06_detail_crop.png','Figure 6. Identical image crop at the original resolution and at ranks 1, 4 and 20.',2.15)
p('Suit seams, the badge and equipment boundaries become indistinct before broad orange regions disappear. This supports the observed loss of fine detail at low rank in this image. Small singular values are not universally equivalent to high spatial frequencies, so the observation is not asserted as a theorem for every image.')
p('These answers are based on the generated figures and measured metrics. No classroom discussion attendance, group contribution or human observation session is fabricated.')

page('8. Eigenfaces: method and evaluation design')
p('Eigenfaces uses a shared PCA basis for both compression and recognition [5]. Each image is flattened to D=10304 entries. A training-only mean is subtracted before SVD of the centered training matrix. The leading rows of V-transpose are the shared basis B_k.')
eq('z = (x−μ)Bₖᵀ;  x̂ = μ+zBₖ',r'z=(x-\mu)B_k^T,\qquad \widehat{x}=\mu+zB_k.')
p('Nearest-neighbor classification compares Euclidean distances between coefficient vectors and returns the label of the closest training photograph. Per subject, a seeded permutation chooses seven training and three test photographs: 280 training and 120 test images overall. Seed 20260917 is retained for reproducibility. Training and test filenames never overlap. The mean, basis and quantization parameters are fitted only on training images.')
p('The principal evaluation uses k=80, fixed before evaluating this run. Other ranks are sensitivity results, not a test-set search for a new optimum. This is closed-set recognition: the identities are present in training, but their test photographs are withheld. It does not measure recognition of previously unseen identities.')
fig('08_mean_eigenfaces.png','Figure 7. Training mean and the first eleven eigenfaces, each visualized on its own grayscale range.',3.45)
p(f'The centered training matrix has numerical rank {f["centered_rank"]}. The basis orthogonality residual is {f["orthogonality_error"]:.3e}. Eigenfaces are signed spatial patterns, not additional face photographs. The recorded split and all k=80 predictions are included as CSV files.')

page('9. Eigenfaces: recognition and reconstruction results')
tab(['k','Correct /120','Accuracy (%)','Test PSNR (dB)','Library param. CR'],[[r['k'],r['correct'],f'{r["accuracy"]*100:.2f}',f'{r["test_psnr_db"]:.3f}',f'{r["train_library_parameter_ratio"]:.3f}'] for r in f['metrics']],8.5)
p(f'The primary result is {f["primary_correct"]}/120 correct, or {f["primary_accuracy"]*100:.2f}%. The descriptive sample-level Wilson 95% interval is [{f["wilson_95"][0]*100:.2f}%, {f["wilson_95"][1]*100:.2f}%]. Photographs from a subject are related, and a single random split does not establish cross-identity or real-world generalization.')
fig('09_face_reconstructions.png','Figure 8. Three held-out test photographs and their shared-basis reconstructions.',3.25)
eq('Shared-library parameter CR = ND/[D+kD+Nk]',r'\mathrm{CR}_{\mathrm{library}}=\frac{ND}{D+kD+Nk}.')
p('The denominator includes the mean, basis and all N coefficient vectors. For N=280 and k=80 the scalar ratio is 3.366. The byte ratio additionally depends on dtype. At k=200 and 279, reconstruction PSNR improves while recognition accuracy falls to 96.67%; better reconstruction need not imply better classification. A held-out image may contain components outside the training subspace, so keeping every nonzero training component still does not reconstruct it exactly.')

page('10. Shared-basis upload and int8 coefficient quantization')
p('This extension is a local encoder/decoder experiment on real ORL images, not an edge-device deployment. It uses k=32. For each dimension, the 0.1% and 99.9% training-coefficient quantiles define lo and hi. Quantization parameters are not fitted to test images.')
eq('scale = (hi−lo)/255;  q = round(clip((z−lo)/scale,0,255))−128',r's=\frac{\mathrm{hi}-\mathrm{lo}}{255},\qquad q=\operatorname{round}\!\left(\operatorname{clip}\left(\frac{z-\mathrm{lo}}s,0,255\right)\right)-128.')
eq('ẑ = (q+128)scale+lo;  x̂ = μ+ẑB₃₂',r'\widehat{z}=(q+128)s+\mathrm{lo},\qquad \widehat{x}=\mu+\widehat{z}B_{32}.')
tab(['Measurement','Observed result'],[
['Float coefficient reconstruction PSNR',f'{q["float_psnr_db"]:.6f} dB'],
['Int8 coefficient reconstruction PSNR',f'{q["int8_psnr_db"]:.6f} dB'],
['PSNR loss',f'{q["psnr_loss_db"]:.6f} dB'],
['Maximum pixel change from float reconstruction',f'{q["max_pixel_difference"]:.6f} on [0,1]'],
['Coefficients outside training quantile limits',f'{q["clipped_coefficient_fraction"]*100:.4f}%'],
['Nearest-neighbor accuracy after quantizing both sets',f'{q["int8_recognition_accuracy"]*100:.2f}%'],
['Raw uint8 frame / int8 coefficients','10,304 B / 32 B'],
['Shared model raw bytes / NPZ file',f'{q["shared_model_raw_bytes"]:,} B / {q["shared_model_npz_bytes"]:,} B'],
['Ratio including shared model, first 120 frames',f'{q["first_120_frames_ratio"]:.4f}'],
['First frame count with total raw-byte savings',q['break_even_frames']]],9)
p('The model includes the float32 basis, mean, lo and scale. Its SHA-256 identifier and version are saved with the coefficients. Loading the saved model and decoding the actual int8 bytes is numerically cross-checked. Both endpoints must agree on the basis ID, version and normalization; unknown versions require model synchronization. ORL inputs already use fixed dimensions and division by 255, so no per-frame brightness standardization metadata is used here.')
p('The 322-fold per-frame payload ratio excludes the shared model. Including the model needs at least 133 frames to save raw pixel bytes; packet headers, retransmission and security metadata are excluded. The small mean PSNR loss does not eliminate local distortion: the maximum pixel change is 0.021111, not the handout example of 0.0015.')
p('Projection is not encryption or proven anonymization. These coefficients still allow identity recognition, and a party with the basis can approximately reconstruct faces. Discarding leading components as a privacy measure was not evaluated.')

page('11. Three-dimensional shared basis: a teaching simulation')
p('The original 3D Morphable Model represents aligned three-dimensional shape and texture using separate learned means and PCA bases [6]. Its shape representation can be written as a mean shape plus a weighted sum of shape modes; texture uses an analogous model.')
eq('S = μS + Σᵢ αᵢbᵢ;  T = μT + Σᵢ βᵢcᵢ',r'S=\mu_S+\sum_i\alpha_i b_i,\qquad T=\mu_T+\sum_i\beta_i c_i.')
p('No licensed real 3D face scans were supplied. To demonstrate the shared-basis mechanism, a synthetic surface uses a 40 x 32 grid (1280 vertices, 3840 coordinate entries) and six linear deformation modes. A fixed seed generates 200 surfaces, with 160 for training and 40 for testing. This is a controlled simulation, not a scanned human face or a fitted real 3DMM. Texture PCA and photo-to-3D fitting are not implemented.')
fig('11_synthetic_3dmm.png','Figure 9. Synthetic surface and low-rank reconstruction; no real human scan is shown.',3.1)
tab(['Retained modes k','Held-out coordinate RMSE'],[[r['k'],f'{r["test_rmse"]:.6e}'] for r in d['metrics']])
p('At k=6 the reconstruction reaches floating-point precision because the data generator has exactly six linear deformation degrees of freedom. This validates the code under its construction assumptions; it does not imply that real three-dimensional faces need only six coefficients. A true 3DMM experiment additionally requires licensed scans, consistent vertex correspondence, and independent evaluation on real faces.')

page('12. Temporal shared basis: a teaching simulation')
p('ORL is a static photograph database; filename order is not a measured time sequence. The temporal extension therefore explicitly generates 300 frames of 32-dimensional coefficient trajectories. A held-out face coefficient vector is perturbed by smooth sinusoids with amplitude 0.12 times each training coefficient standard deviation. The training-fitted quantizer is reused.')
fig('10_synthetic_temporal_coefficients.png','Figure 10. The first four synthetic coefficient trajectories; these are not measured video frames.',2.5)
p('An absolute coefficient reset is inserted every 60 frames; other frames store differences from the preceding quantized frame. Differences are stored as int16 because they may exceed the int8 range. zlib level 9 is applied to both direct and differential representations for a fair comparison. Decoding recovers every quantized coefficient exactly, and compressed bytes pass a round-trip check.')
tab(['Representation','Actual bytes'],[['Direct int8 coefficients',t['direct_int8_bytes']],['Uncompressed int16 differences',t['delta_int16_bytes']],['Direct int8 coefficients + zlib',t['direct_zlib_bytes']],['Int16 differences + zlib',t['delta_zlib_bytes']]])
p(f'Differential coding gives an additional factor of {t["extra_factor_vs_direct_zlib"]:.3f} relative to directly zlib-compressed coefficients: 1979 B versus 1387 B. This does not reproduce the handout\'s illustrative factor of 2 to 3. Without entropy compression, int16 differencing doubles the coefficient payload. Including one shared-model transfer, the 300-frame ratio relative to raw uint8 frames is {t["ratio_with_model_vs_raw_pixels"]:.3f}.')
p('Lossless recovery applies to the already quantized coefficients, not the original image pixels. The smoothness was deliberately built into the simulation. Absolute resets restore predictor state but do not retrain the PCA basis. Scene-change detection, pose-drift detection, basis updates and comparisons against JPEG or video codecs remain untested; those mechanisms are design proposals rather than measured production results.')

page('13. Analysis, conclusions and reproducibility')
h('13.1 Conclusions')
p('The supplied matrix has rank 4 with singular values spanning five orders of magnitude. Two terms retain nearly all energy, while all four are required for exact reconstruction. The left/right eigenvector identities and the rank-one sum agree with numerical reference results.')
p('For images, SVD produces a controlled fidelity-storage tradeoff. Energy, relative error, PSNR, scalar count and actual byte count describe different aspects and must be reported separately. The 512 x 512 image reaches a scalar break-even rank of 255.750, but float32 factors exceed the raw uint8 payload from k=64 onward.')
p('A shared basis unifies face reconstruction and closed-set recognition. The measured k=80 result is 118/120, or 98.33%, for one reproducible split. A training-only quantizer permits small int8 coefficient payloads, but model transfer and local clipping errors must be included. Synthetic 3D and temporal experiments validate mechanisms under explicit construction assumptions.')
h('13.2 Actual environment and execution')
tab(['Item','Recorded value'],[['Execution date',run['run_date']],['Operating system',run['platform']],['Python',run['python']],['NumPy / SciPy',run['numpy']+' / '+run['scipy']],['Matplotlib / Pillow',run['matplotlib']+' / '+run['pillow']],['Seed',run['seed']],['Numerical checks',f'{run["checks_passed"]}/{run["checks_total"]} passed'],['Core computation elapsed time',f'{run["elapsed_seconds"]:.3f} s']],9)
p('Elapsed time excludes data acquisition, dependency installation and document generation. Codex assisted with requirement extraction, Python implementation, execution and report production. Results are based on executed code and recorded outputs; personal details and assessment fields are deliberately blank. No human group-work record is invented.')
h('13.3 Reproduction')
code('python -m venv .venv\nsource .venv/bin/activate\npython -m pip install -r requirements-report.txt\npython 02_实验本体/run_experiment.py\npython build_reports.py\nbash 04_LaTeX源码/compile.sh')
p('Run these commands from the package root. The data are included, so numerical reruns need no download. The saved lock file records exact Python-package versions. Export the generated DOCX using Microsoft Word or LibreOffice to refresh its PDF; this delivery uses LibreOffice. Data hashes, numerical checks and final PDF validation are provided in the package. Changes to personal details do not require rerunning the experiment.')

page('Appendix A. Original numeric inputs and additional checks')
mat('Q5 (5 x 5), transcribed row by row from the course PDF',npz['Q5'])
mat('Q4 (4 x 4), acting as V-transpose in the supplied product',npz['Q4'])
eq('Σ is 5×4: diag(100,10,0.1,0.001), followed by a zero row',r'\Sigma=\begin{pmatrix}100&0&0&0\\0&10&0&0\\0&0&0.1&0\\0&0&0&0.001\\0&0&0&0\end{pmatrix}.')
mat('A as printed in the course PDF',npz['A_printed'])
p('Direct SVD of the printed, rounded matrix gives singular values '+', '.join(f'{x:.9f}' for x in m['printed_A_singular_values'])+'. The report consistently uses the assembled version; both arrays are retained so that rounding and assembly conventions are traceable.')
p('The verification record contains 22 passed assertions, including matrix agreement with the printed entries, matrix rank and reconstruction, left/right eigenvector equations, the teaching implementation, seven image-tail identities, monotonic error, full-rank zero tail, all 400 face files and their dimensions, disjoint train/test splits, saved-model decoding, differential decoding, compressed-byte recovery, and six-mode synthetic reconstruction.')

page('References and evidence scope')
refs=[
'[1] Course handout. First Computer Laboratory: SVD and Its Applications. Five-page PDF supplied by the user. The original report template was also supplied by the user; both originals are preserved in 01_Source_Materials.',
'[2] AT&T Laboratories Cambridge. The Database of Faces (formerly ORL). https://www.cl.cam.ac.uk/research/dtg/attarchive/facedatabase.html. Dataset: 40 subjects, 10 photographs each, 92 x 112 pixels. Source page checked on 9 October 2026; archived input obtained on 17 September 2026.',
'[3] scikit-image. skimage.data.astronaut. https://scikit-image.org/docs/stable/api/skimage.data.html#skimage.data.astronaut. NASA image credit. This experiment uses the archived v0.25.2 image, not a dynamically changed input. Source page checked on 9 October 2026.',
'[4] NumPy. numpy.linalg.svd documentation. https://numpy.org/doc/stable/reference/generated/numpy.linalg.svd.html. Consulted for reduced shapes, reconstruction and singular-vector relations on 9 October 2026. Runtime versions are recorded separately.',
'[5] Turk, M. and Pentland, A. (1991). Eigenfaces for Recognition. Journal of Cognitive Neuroscience, 3(1), 71-86. DOI: 10.1162/jocn.1991.3.1.71. A saved paper copy is included for method background.',
'[6] Blanz, V. and Vetter, T. (1999). A Morphable Model for the Synthesis of 3D Faces. Proceedings of SIGGRAPH 1999, 187-194. DOI: 10.1145/311535.311556. Author project archive: https://shapemodelling.cs.unibas.ch/gravis-site-archive/Sigg99.html. Checked on 9 October 2026; saved paper included.'
]
for s in refs:p(s)
h('What was measured')
p('Matrix SVD, image metrics and reconstructions, ORL Eigenfaces classification, local int8 coefficient decoding, model byte counts, and the controlled synthetic 3D/temporal examples were executed. Values in this report are read from the current results JSON/CSV files. Figures are generated directly by the same run.')
h('What was not established')
p('The experiment does not establish real-device latency, actual network transfer, real 3D-face reconstruction accuracy, real-video coding gains, scene-change robustness, privacy protection, or generalization to unseen identities. The handout\'s citation counts and illustrative performance numbers are not reported as current measurements.')

# Word: begin with the actual supplied template, translate labels and fill its cells.
doc=Document(ROOT/'01_指导资料/Original_Report_Template.docx')
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.03
for name in ['Title','Heading 1','Heading 2']:
    if name not in doc.styles:
        doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    doc.styles[name].font.name='Times New Roman'
    doc.styles[name].paragraph_format.line_spacing=1.0
    doc.styles[name].paragraph_format.space_before=Pt(4)
    doc.styles[name].paragraph_format.space_after=Pt(4)
doc.styles['Heading 2'].base_style=normal
doc.styles['Heading 2'].paragraph_format.keep_with_next=True
doc.core_properties.title='SVD and Its Applications - English Laboratory Report'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
doc.core_properties.subject='Numerical Analysis Computer Laboratory'
doc.paragraphs[0].text='Beihang University\nSchool of Mathematical Sciences\nLaboratory Report'
doc.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in doc.paragraphs[0].runs:r.font.name='Times New Roman';r.font.size=Pt(15);r.bold=True
doc.paragraphs[1].text='Detailed experimental procedures, results and analysis continue on the following pages.'
for r in doc.paragraphs[1].runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=False
T=doc.tables[0]
values={
(0,0):'Course: Numerical Analysis\nComputer Laboratory',
(0,2):'Experiment: SVD and Its Applications',
(1,0):'Experiment type: Demonstration [ ]  Verification [ ]\nComprehensive [X]  Design [ ]',
(2,0):'Class: 230116', (2,1):'Name: yule yang', (2,3):'Student ID: 23376268',
(3,0):'Lab date: __________',(3,1):'Instructor: __________',(3,3):'Grade: __________',
(4,0):f'Experimental environment (hardware and software):\nmacOS on Apple Silicon; Python {run["python"]}; NumPy {run["numpy"]}; SciPy {run["scipy"]}; Matplotlib {run["matplotlib"]}; Pillow {run["pillow"]}. AI assistant: Codex. Numerical execution: 9 October 2026.',
(5,0):'Objectives and experimental content:\nAnalyze the supplied matrix using SVD, singular vectors, rank-one terms and retained energy. Complete grayscale and RGB low-rank reconstruction, all-rank metrics, diagnostic curves and the three key findings. Investigate Eigenfaces and int8 shared-basis coding; demonstrate 3D and temporal extensions using explicitly labeled simulations. Full details: Sections 1-12 and Appendix A.',
(6,0):'Experimental procedure and results:\nAssemble A = Q5 Sigma Q4 and cross-check against the printed matrix. Compute rank 4 and singular values approximately 100, 10, 0.1 and 0.001. Evaluate RGB ranks 1-512 and generate the required plots. Parameter break-even rank: 255.750. Train/test Eigenfaces on 280/120 ORL images; k=80 gives 118/120 correct (98.33%). Int8 coefficients use 32 B per frame; total savings begin at 133 frames after counting the model. All 22 numerical checks passed. Detailed tables, equations, figures and code follow on continuation pages.',
(7,0):'Experimental analysis and summary:\nLow-rank SVD balances fidelity and storage. High retained energy can coexist with poor visual detail; parameter ratios differ from byte ratios. Shared bases support recognition and compression but incur model-transfer costs. Synthetic 3D and time-series results are mechanism demonstrations. Section 13 discusses conclusions and reproduction; the final page records evidence limits.'
}
for (i,j),s in values.items():
    c=T.cell(i,j);c.text=s
    for para in c.paragraphs:
        para.paragraph_format.space_before=Pt(1);para.paragraph_format.space_after=Pt(3)
        para.paragraph_format.line_spacing=1.0
        for r in para.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
for row in T.rows:
    if row._tr.trPr is not None:
        for el in list(row._tr.trPr):
            if el.tag in [qn('w:trHeight'),qn('w:cantSplit')]:row._tr.trPr.remove(el)
    pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
section=doc.add_section(WD_SECTION_START.NEW_PAGE)
section.top_margin=Mm(19);section.bottom_margin=Mm(19);section.left_margin=Mm(20);section.right_margin=Mm(20)
section.header_distance=Mm(8);section.footer_distance=Mm(8)
section.header.is_linked_to_previous=False
hp=section.header.paragraphs[0];hp.text='SVD and Its Applications | Detailed Laboratory Report'
hp.runs[0].font.size=Pt(9);hp.runs[0].font.name='Times New Roman'
section.footer.is_linked_to_previous=False
for sec in doc.sections:
    for el in list(sec._sectPr):
        if el.tag==qn('w:docGrid'):sec._sectPr.remove(el)
for el in doc.settings._element.iter(qn('w:compatSetting')):
    if el.get(qn('w:name'))=='compatibilityMode':el.set(qn('w:val'),'15')
fp=section.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
for idx,pg in enumerate(pages):
    ph=doc.add_paragraph(pg['title'],'Heading 1');ph.paragraph_format.space_after=Pt(8)
    ph.paragraph_format.page_break_before=idx>0
    ph.runs[0].font.size=Pt(15);ph.runs[0].font.color.rgb=RGBColor.from_string('17365D')
    for b in pg['blocks']:
        kind=b['kind']
        if kind=='paragraph':doc.add_paragraph(b['text'])
        elif kind=='heading':
            ph=doc.add_paragraph(b['text'],'Heading 2');ph.runs[0].font.size=Pt(11.5)
            ph.paragraph_format.space_before=Pt(6);ph.paragraph_format.space_after=Pt(3)
        elif kind=='equation':
            para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after=Pt(6)
            para._p.append(office_math(b['latex']))
        elif kind=='code':
            para=doc.add_paragraph();para.paragraph_format.space_after=Pt(6)
            for n,line in enumerate(b['text'].splitlines()):
                if n:para.add_run().add_break()
                rr=para.add_run(line);rr.font.name='Courier New';rr.font.size=Pt(8.5)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F4F7');para._p.get_or_add_pPr().append(sh)
        elif kind=='table':
            tb=doc.add_table(rows=1,cols=len(b['headers']));tb.style='Table Grid'
            for c,s in zip(tb.rows[0].cells,b['headers']):c.text=s
            head=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(head)
            for r in b['rows']:
                for c,s in zip(tb.add_row().cells,r):c.text=s
            for ri,row in enumerate(tb.rows):
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for c in row.cells:
                    for para in c.paragraphs:
                        para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1
                        for rr in para.runs:rr.font.name='Times New Roman';rr.font.size=Pt(b['size']);rr.bold=ri==0
                    if ri==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF4');c._tc.get_or_add_tcPr().append(sh)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif kind=='figure':
            from PIL import Image
            w,h0=Image.open(OUT/b['name']).size
            height=min(b['height'],6.65*h0/w)
            para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
            para.add_run().add_picture(str(OUT/b['name']),height=Inches(height))
            cp=doc.add_paragraph(b['caption']);cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_after=Pt(6)
            for rr in cp.runs:rr.font.size=Pt(9);rr.italic=True
doc.save(ROOT/'03_最终实验报告/SVD_Report_English_Word.docx')
math_cache_path.write_text(json.dumps({'version':2,'equations':math_cache},ensure_ascii=False,indent=2))

# LaTeX: identical content blocks, with native mathematical typesetting.
def esc(s):
    mp={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}'}
    return ''.join(mp.get(c,c) for c in str(s))
tex=[r'''\documentclass[a4paper,11pt]{article}
\usepackage[margin=19mm,headheight=14pt]{geometry}
\usepackage{fontspec}
\setmainfont{Times New Roman}[Path=/System/Library/Fonts/Supplemental/,Extension=.ttf,UprightFont=*,BoldFont=* Bold,ItalicFont=* Italic,BoldItalicFont=* Bold Italic]
\setmonofont{Courier New}[Path=/System/Library/Fonts/Supplemental/,Extension=.ttf,UprightFont=*,BoldFont=* Bold,ItalicFont=* Italic,BoldItalicFont=* Bold Italic]
\usepackage{amsmath,amssymb,booktabs,graphicx,float,tabularx,array,xcolor,fancyhdr,fvextra,hyperref,xurl}
\hypersetup{colorlinks=true,linkcolor=blue!40!black,urlcolor=blue!40!black,pdftitle={SVD and Its Applications - English Laboratory Report}}
\setlength{\parindent}{0pt}
\setlength{\parskip}{5pt}
\renewcommand{\arraystretch}{1.16}
\pagestyle{fancy}\fancyhf{}\fancyhead[L]{SVD and Its Applications}\fancyhead[R]{Laboratory Report}\fancyfoot[C]{\thepage}
\DefineVerbatimEnvironment{code}{Verbatim}{fontsize=\small,breaklines=true,frame=single}
\begin{document}
\thispagestyle{empty}
\begin{center}
{\Large Beihang University\\School of Mathematical Sciences\par}
\vspace{10mm}
{\LARGE\bfseries SVD and Its Applications\par}
\vspace{4mm}{\Large English Laboratory Report\par}
\vspace{8mm}
\end{center}
\begin{tabularx}{\linewidth}{|X|X|}\hline
Course: Numerical Analysis\newline Computer Laboratory & Experiment: SVD and Its Applications\\\hline
\multicolumn{2}{|l|}{Experiment type: Comprehensive}\\\hline
Class: \underline{\hspace{22mm}} & Name: \underline{\hspace{25mm}}\\\hline
Student ID: \underline{\hspace{22mm}} & Lab date: \underline{\hspace{23mm}}\\\hline
Instructor: \underline{\hspace{22mm}} & Grade: \underline{\hspace{26mm}}\\\hline
\end{tabularx}
\vspace{6mm}
\textbf{Objectives and experimental content}

Analyze the supplied matrix using SVD, singular vectors, rank-one terms and retained energy. Complete grayscale and RGB low-rank reconstruction, all-rank metrics, diagnostic curves and the three key findings. Investigate Eigenfaces and int8 shared-basis coding; demonstrate 3D and temporal extensions using explicitly labeled simulations.

\textbf{Experimental procedure and results}

The assembled matrix has rank 4 and singular values approximately 100, 10, 0.1 and 0.001. RGB image metrics cover every rank from 1 to 512. The parameter break-even rank is 255.750. Eigenfaces uses 280 training and 120 test images; k=80 gives 118/120 correct (98.33\%). Int8 coefficients use 32 B per frame; total savings begin at 133 frames after counting the shared model. All 22 numerical checks passed.

\textbf{Experimental analysis and summary}

Low-rank SVD balances fidelity and storage. High retained energy can coexist with poor visual detail; parameter ratios differ from byte ratios. Shared bases support recognition and compression but incur model-transfer costs. Synthetic 3D and temporal results are mechanism demonstrations.

\textbf{Environment}

macOS on Apple Silicon; Python, NumPy, SciPy, Matplotlib and Pillow. Codex assisted with implementation and report generation. Exact versions and the execution date are recorded in Section 13. Personal and assessment fields remain blank.
''']
ld=ROOT/'04_LaTeX源码';(ld/'figures').mkdir(exist_ok=True)
for pg in pages:
    tex += [r'\clearpage',r'\section*{'+esc(pg['title'])+'}']
    for b in pg['blocks']:
        k=b['kind']
        if k=='paragraph':
            chunks=re.split(r'(https?://\S+)',b['text'])
            result=[]
            for s in chunks:
                if s.startswith(('https://','http://')):
                    suffix='.' if s.endswith('.') else ''
                    if suffix:s=s[:-1]
                    result.append(r'\url{'+s+'}'+suffix)
                else:result.append(esc(s))
            tex.append(''.join(result)+'\n')
        elif k=='heading':tex.append(r'\subsection*{'+esc(b['text'])+'}')
        elif k=='equation':tex.append(r'\['+b['latex']+r'\]')
        elif k=='code':tex.append(r'\begin{code}'+'\n'+b['text']+'\n'+r'\end{code}')
        elif k=='figure':
            shutil.copy2(OUT/b['name'],ld/'figures'/b['name'])
            tex.append(r'\begin{figure}[H]\centering\includegraphics[width=.97\linewidth,height='+str(b['height'])+r'in,keepaspectratio]{figures/'+b['name']+r'}\caption{'+esc(b['caption'].split('. ',1)[1])+r'}\end{figure}')
        elif k=='table':
            n=len(b['headers']);spec='@{}'+('l' if n>2 else r'>{\raggedright\arraybackslash}X')*n+'@{}'
            if n>2:
                tex.append(r'\begin{center}\fontsize{'+str(b['size'])+'}{'+str(b['size']+2)+r'}\selectfont\begin{tabular}{'+spec+r'}\toprule')
            else:tex.append(r'\begin{center}\fontsize{'+str(b['size'])+'}{'+str(b['size']+2)+r'}\selectfont\begin{tabularx}{\linewidth}{'+spec+r'}\toprule')
            tex.append(' & '.join(esc(x) for x in b['headers'])+r'\\\midrule')
            for row in b['rows']:tex.append(' & '.join(esc(x) for x in row)+r'\\')
            tex.append(r'\bottomrule\end{'+('tabular' if n>2 else 'tabularx')+r'}\end{center}')
tex.append(r'\end{document}')
(ld/'main.tex').write_text('\n'.join(tex))
(ROOT/'report_content.json').write_text(json.dumps(pages,indent=2,ensure_ascii=False))
print('Generated Word and LaTeX sources from',run['checks_passed'],'passed checks.')
